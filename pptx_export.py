"""PowerPoint export of an experiment: a 16:9 deck built from the same data as
the PDF report.

Slides: title (logo, details), summary, data files, metadata (the compact
layout of ``metasummary.layout_file`` as native tables), the camera pictures
and SnapMaps (one picture per slide, see ``imagepages``), then one slide per
figure page (a picture, an editable caption, speaker notes describing the
look). Needs python-pptx; no Tk here.
"""

from __future__ import annotations

import datetime
import io
import os

import appinfo
import covers
import metasummary
import panelview
import reportspec

SECTIONS = ("title", "files", "metadata", "images", "figures")

SLIDE_W, SLIDE_H = 13.333, 7.5          # inches (16:9)
MARGIN = 0.6
BODY_W = SLIDE_W - 2 * MARGIN
FIGURE_SIZE = (12.1, 4.95)              # inches, for the picture on a figure slide
NAVY = (0x2C, 0x3E, 0x50)
GREY = (0x6B, 0x77, 0x85)
INK = (0x1A, 0x21, 0x27)
ALT = (0xF2, 0xF5, 0xF8)
FONT = "Calibri"

ROW_H = 0.29                            # table row height (10 pt text)
TABLE_TOP = 1.25
BOTTOM = SLIDE_H - 0.55                 # keep clear of the footer


class PptxError(Exception):
    """The deck could not be built (message is user-facing)."""


def _modules():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        raise PptxError("PowerPoint export needs python-pptx "
                        "(pip install python-pptx).")
    return (Presentation, Inches, Pt, RGBColor, MSO_ANCHOR, PP_ALIGN,
            MSO_SHAPE)


# -- pure helpers ----------------------------------------------------------------
def paragraphs(text):
    """Blank-line separated paragraphs of free text (single newlines kept)."""
    return [b.strip("\n") for b in (text or "").replace("\r\n", "\n")
            .split("\n\n") if b.strip()]


def chunk_paragraphs(paras, max_lines=11, chars_per_line=92):
    """Split paragraphs over slides so each holds about ``max_lines`` lines
    (long paragraphs are cut at word boundaries)."""
    def lines(p):
        return sum(max(1, -(-len(part) // chars_per_line))
                   for part in p.split("\n")) + 1     # + spacing after

    pieces = []
    for p in paras:
        if lines(p) <= max_lines:
            pieces.append(p)
            continue
        cur = ""
        cap = (max_lines - 1) * chars_per_line
        words = []
        for word in p.split(" "):          # an unbreakable run is cut too
            while len(word) > cap:
                words.append(word[:cap])
                word = word[cap:]
            words.append(word)
        for word in words:
            if lines(cur + " " + word) > max_lines and cur:
                pieces.append(cur)
                cur = word
            else:
                cur = (cur + " " + word) if cur else word
        if cur:
            pieces.append(cur)
    slides, cur, used = [], [], 0
    for p in pieces:
        n = lines(p)
        if cur and used + n > max_lines:
            slides.append(cur)
            cur, used = [], 0
        cur.append(p)
        used += n
    if cur:
        slides.append(cur)
    return slides


def look_notes(state):
    """Speaker-notes text describing a saved look."""
    st = state or {}
    parts = [f"View: {st.get('view_mode', 'Stack')}",
             f"grouped by {st.get('group_by', 'element name').lower()}",
             f"{st.get('energy_scale', 'Binding').lower()} energy axis"]
    if st.get("norm") not in (None, "None"):
        parts.append(f"normalised: {st['norm']}")
    if st.get("colour_scale") not in (None, "Theme default"):
        parts.append(f"colour scale: {st['colour_scale']}"
                     + (" (reversed)" if st.get("colour_reverse") else ""))
    if st.get("view_mode") in ("Waterfall 3D", "Heatmap"):
        parts.append(f"z axis: {st.get('z_axis', 'Auto')}")
    ticked = st.get("ticked") or []
    names = metasummary.compact_labels(
        [f"{t.get('name', '')}" for t in ticked], limit=12)
    text = "; ".join(parts) + "."
    defaults = {"view": st.get("view_mode", "Stack"),
                "norm": st.get("norm", "None"),
                "offset": st.get("offset", 0.6),
                "z_axis": st.get("z_axis", "Auto"),
                "reverse": bool(st.get("reverse")),
                "fit_show": st.get("fit_show") or {}}
    own = [f"{label}: {panelview.describe(ov, defaults)}"
           for label, ov in panelview.sanitise_all(
               st.get("panel_views")).items()]
    if own:
        text += "\nPanels with their own view: " + "; ".join(own) + "."
    if ticked:
        text += f"\nSpectra shown ({len(ticked)}): {names}."
    return text


def pack_items(items, budget):
    """Greedy page packing: ``items`` is ``[(height, payload)]``; returns
    ``[[payload, ...], ...]`` with each page's total height <= ``budget``
    (an oversize item gets a page to itself)."""
    pages, cur, used = [], [], 0.0
    for h, payload in items:
        if cur and used + h > budget:
            pages.append(cur)
            cur, used = [], 0.0
        cur.append(payload)
        used += h
    if cur:
        pages.append(cur)
    return pages


# -- drawing -----------------------------------------------------------------------
class _Deck:
    def __init__(self, title):
        (self.Presentation, self.Inches, self.Pt, self.RGB, self.ANCHOR,
         self.ALIGN, self.SHAPE) = _modules()
        self.prs = self.Presentation()
        self.prs.slide_width = self.Inches(SLIDE_W)
        self.prs.slide_height = self.Inches(SLIDE_H)
        self.title = title

    # low level
    def rgb(self, t):
        return self.RGB(*t)

    def bar(self, slide):
        shp = slide.shapes.add_shape(self.SHAPE.RECTANGLE, 0, 0,
                                     self.Inches(SLIDE_W), self.Inches(0.14))
        shp.fill.solid()
        shp.fill.fore_color.rgb = self.rgb(NAVY)
        shp.line.fill.background()

    def footer(self, slide, n):
        self.text(slide, MARGIN, SLIDE_H - 0.42, BODY_W, 0.3,
                  [f"{self.title}   |   {n}"], size=9, color=GREY)

    def text(self, slide, x, y, w, h, paras, size=14, bold=False,
             color=INK, align=None, space_after=6):
        box = slide.shapes.add_textbox(self.Inches(x), self.Inches(y),
                                       self.Inches(w), self.Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = self.Pt(space_after)
            if align:
                p.alignment = align
            r = p.add_run()
            r.text = para
            r.font.size = self.Pt(size)
            r.font.bold = bold
            r.font.name = FONT
            r.font.color.rgb = self.rgb(color)
        return box

    def content_slide(self, title, number):
        """A 'Title Only' slide with the title restyled to fit 16:9."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        self.bar(slide)
        t = slide.shapes.title
        t.left, t.top = self.Inches(MARGIN), self.Inches(0.35)
        t.width, t.height = self.Inches(BODY_W), self.Inches(0.75)
        tf = t.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = self.ANCHOR.MIDDLE
        tf.text = title
        p = tf.paragraphs[0]
        p.alignment = self.ALIGN.LEFT
        for r in p.runs:
            r.font.size = self.Pt(26 if len(title) <= 52 else 20)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = self.rgb(NAVY)
        self.footer(slide, number)
        return slide

    def table(self, slide, x, y, w, weights, header, rows, size=10,
              row_h=ROW_H, header_fill=NAVY):
        nrows = len(rows) + (1 if header else 0)
        shape = slide.shapes.add_table(nrows, len(weights), self.Inches(x),
                                       self.Inches(y), self.Inches(w),
                                       self.Inches(row_h * nrows))
        tbl = shape.table
        tbl.horz_banding = False
        total = float(sum(weights))
        widths = [self.Inches(w * v / total) for v in weights]
        widths[-1] = self.Inches(w) - sum(widths[:-1])     # exact total
        for col, wd in zip(tbl.columns, widths):
            col.width = wd
        for r in tbl.rows:
            r.height = self.Inches(row_h)
        allrows = ([list(header)] if header else []) + rows
        for ri, row in enumerate(allrows):
            is_head = bool(header) and ri == 0
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.margin_left = cell.margin_right = self.Inches(0.05)
                cell.margin_top = cell.margin_bottom = self.Inches(0.02)
                cell.vertical_anchor = self.ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb(
                    header_fill if is_head else
                    (ALT if (ri % 2 == 0) else (255, 255, 255)))
                tf = cell.text_frame
                tf.word_wrap = True
                tf.text = str(val)
                for p in tf.paragraphs:
                    for r in p.runs:
                        r.font.size = self.Pt(size)
                        r.font.name = FONT
                        r.font.bold = is_head
                        r.font.color.rgb = self.rgb((255, 255, 255)
                                                    if is_head else INK)
        return shape


# -- slide builders ------------------------------------------------------------------
def _title_slide(deck, details, logo, art=None):
    slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[0])
    deck.bar(slide)
    if art is not None and art.data:           # a strip along the bottom
        slide.shapes.add_picture(
            io.BytesIO(art.data), 0, deck.Inches(SLIDE_H - art.height),
            width=deck.Inches(SLIDE_W))
    ph = list(slide.placeholders)
    title = (details.get("title") or "").strip() or "Experiment report"
    t = slide.shapes.title
    t.left, t.top = deck.Inches(0.8), deck.Inches(2.3)
    t.width, t.height = deck.Inches(SLIDE_W - 1.6), deck.Inches(1.5)
    t.text_frame.word_wrap = True
    t.text_frame.vertical_anchor = deck.ANCHOR.BOTTOM
    t.text_frame.text = title
    for p in t.text_frame.paragraphs:
        p.alignment = deck.ALIGN.LEFT
        for r in p.runs:
            r.font.size = deck.Pt(38)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = deck.rgb(NAVY)
    date = (details.get("date") or "").strip() \
        or datetime.date.today().isoformat()
    lines = [f"{label}: {v}" for label, v in
             (("Customer", details.get("customer")),
              ("Reference", details.get("reference")),
              ("Operator", details.get("operator")), ("Date", date))
             if v and str(v).strip()]
    sub = [p for p in ph if p.placeholder_format.idx == 1]
    if sub:
        s = sub[0]
        s.left, s.top = deck.Inches(0.8), deck.Inches(4.0)
        s.width, s.height = deck.Inches(SLIDE_W - 1.6), deck.Inches(1.8)
        tf = s.text_frame
        tf.word_wrap = True
        tf.text = lines[0] if lines else ""
        for line in lines[1:]:
            tf.add_paragraph().text = line
        for p in tf.paragraphs:
            p.alignment = deck.ALIGN.LEFT
            p.space_after = deck.Pt(4)
            for r in p.runs:
                r.font.size = deck.Pt(20)
                r.font.name = FONT
                r.font.color.rgb = deck.rgb(GREY)
    if logo and os.path.isfile(logo):
        try:
            pic = slide.shapes.add_picture(logo, deck.Inches(0.8),
                                           deck.Inches(0.7),
                                           width=deck.Inches(3.2))
            if pic.height > deck.Inches(1.3):          # keep it compact
                ratio = deck.Inches(1.3) / pic.height
                pic.height = deck.Inches(1.3)
                pic.width = int(pic.width * ratio)
        except Exception:
            pass                                   # a bad logo must not stop the deck
    return slide


def _text_slides(deck, title, paras, n0, size=16):
    """Slides of a heading and free text (paginated); ``paras`` is a list of
    paragraphs."""
    n = n0
    for i, page in enumerate(chunk_paragraphs(paras)):
        n += 1
        slide = deck.content_slide(title + (" (continued)" if i else ""), n)
        deck.text(slide, MARGIN, 1.3, BODY_W, BOTTOM - 1.3, page,
                  size=size, space_after=10)
    return n


def _files_slides(deck, file_rows, n0, sha="short"):
    n = n0
    per = int((BOTTOM - TABLE_TOP) / ROW_H) - 1
    with_sha = sha != "none"
    for i in range(0, max(1, len(file_rows)), per):
        chunk = file_rows[i:i + per]
        n += 1
        slide = deck.content_slide(
            "Data files" + (" (continued)" if i else ""), n)
        rows = [[r.get("name", ""), r.get("format", ""),
                 str(r.get("regions", "")), _size(r.get("size", 0))]
                + ([(r.get("sha256") or "")[:16]] if with_sha else [])
                for r in chunk]
        deck.table(slide, MARGIN, TABLE_TOP, BODY_W,
                   [4.2, 3.0, 0.9, 1.1, 2.3] if with_sha
                   else [6.4, 4.0, 1.0, 1.5],
                   ["File", "Format", "Regions", "Size"]
                   + (["SHA-256"] if with_sha else []), rows)
    return n


def _size(n):
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024.0


def _metadata_slides(deck, docs, n0, skip=()):
    n = n0
    budget = BOTTOM - TABLE_TOP
    for parser in docs:
        if reportspec.doc_key(parser) in skip:
            continue
        samples = parser.samples_metadata()
        if not samples:
            continue
        name = os.path.basename(parser.path or "experiment")
        lay = metasummary.layout_file(samples)
        items = []                                # (height, draw-callable)
        strips = []

        if lay.common:
            kv = [(metasummary.SHORT.get(k, k), v) for k, v in lay.common]
            rows = [sum(([k, v] for k, v in kv[i:i + 3]), [])
                    for i in range(0, len(kv), 3)]
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            h = 0.32 + ROW_H * len(rows)
            items.append((h, ("common", rows, width)))

        for sl in lay.samples:
            sub = [(metasummary.SHORT.get(k, k), v) for k, v in sl.line]
            line_rows = [sum(([k, v] for k, v in sub[i:i + 3]), [])
                         for i in range(0, len(sub), 3)]
            if line_rows:
                w_ = max(len(r) for r in line_rows)
                line_rows = [r + [""] * (w_ - len(r)) for r in line_rows]
            head_h = 0.36 + ROW_H * len(line_rows)
            max_rows = max(3, int((budget - head_h) / ROW_H) - 1)
            for i in range(0, max(1, len(sl.rows)), max_rows):
                chunk = sl.rows[i:i + max_rows]
                first = i == 0
                h = (head_h if first else 0.36) + ROW_H * (len(chunk) + 1)
                items.append((h, ("sample", sl, chunk, first, line_rows)))
            for title, entries in sl.strips:
                strips.append((f"{sl.name}: {title}", entries))

        for pi, page in enumerate(pack_items(items, budget)):
            n += 1
            slide = deck.content_slide(
                f"Acquisition metadata – {name}"
                + (" (continued)" if pi else ""), n)
            y = TABLE_TOP
            for item in page:
                if item[0] == "common":
                    _, rows, width = item
                    deck.text(slide, MARGIN, y, BODY_W, 0.3,
                              ["Common to every region"], size=12,
                              bold=True, color=NAVY, space_after=0)
                    y += 0.32
                    _kv_table(deck, slide, y, rows, width)
                    y += ROW_H * len(rows)
                else:
                    _, sl, chunk, first, line_rows = item
                    label = sl.name + (
                        f"  ({sl.n_regions} regions)" if first else
                        "  (continued)")
                    deck.text(slide, MARGIN, y, BODY_W, 0.3, [label],
                              size=12, bold=True, color=NAVY, space_after=0)
                    y += 0.36
                    if first and line_rows:
                        _kv_table(deck, slide, y, line_rows,
                                  max(len(r) for r in line_rows))
                        y += ROW_H * len(line_rows)
                    weights = metasummary.column_weights(sl.columns)
                    deck.table(slide, MARGIN, y, BODY_W, weights, sl.columns,
                               [[r.get(c, "") for c in sl.columns]
                                for r in chunk])
                    y += ROW_H * (len(chunk) + 1)

        for title, entries in strips:                # per-level details
            per_slide = 60
            for i in range(0, len(entries), per_slide):
                part = entries[i:i + per_slide]
                n += 1
                slide = deck.content_slide(
                    f"{title} – {name}"
                    + (" (continued)" if i else ""), n)
                percol = -(-len(part) // 3)
                grid = []
                for r in range(percol):
                    row = []
                    for c in range(3):
                        idx = c * percol + r
                        row += list(part[idx]) if idx < len(part) \
                            else ["", "", ""]
                    grid.append(row)
                deck.table(slide, MARGIN, TABLE_TOP, BODY_W,
                           [0.5, 0.8, 2.2] * 3,
                           ["Level", "Etch (s)", "Acquired"] * 3, grid,
                           size=9, row_h=0.26)
    return n


def _kv_table(deck, slide, y, rows, width):
    """Key/value pairs laid out across (label, value, label, value ...)."""
    pairs = width // 2
    weights = [1.0, 1.6] * pairs
    shape = deck.table(slide, MARGIN, y, BODY_W, weights, None, rows,
                       header_fill=NAVY)
    tbl = shape.table
    for ri, row in enumerate(rows):
        for ci in range(0, width, 2):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = deck.rgb((255, 255, 255))
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = deck.rgb(NAVY)
            for cj in (ci + 1,):
                c2 = tbl.cell(ri, cj)
                c2.fill.solid()
                c2.fill.fore_color.rgb = deck.rgb((255, 255, 255))
    return shape


def _figure_slides(deck, figures, render_images, n0):
    n = n0
    for number, fig in enumerate(figures, 1):
        images = render_images(number, fig)
        caption = (fig.get("caption") or "").strip()
        for pi, png in enumerate(images):
            n += 1
            title = f"Figure {number} – {fig.get('name', '')}"
            if pi:
                title += " (continued)"
            slide = deck.content_slide(title, n)
            width, height = FIGURE_SIZE
            pic = slide.shapes.add_picture(
                io.BytesIO(png), deck.Inches((SLIDE_W - width) / 2),
                deck.Inches(1.15), width=deck.Inches(width))
            if pic.height > deck.Inches(height):
                ratio = deck.Inches(height) / pic.height
                pic.height = deck.Inches(height)
                pic.width = int(pic.width * ratio)
                pic.left = int((deck.Inches(SLIDE_W) - pic.width) / 2)
            if caption and pi == 0:
                deck.text(slide, MARGIN, 6.15, BODY_W, 0.9,
                          paragraphs(caption)[:3], size=13, space_after=3)
            notes = slide.notes_slide.notes_text_frame
            notes.text = ((caption + "\n\n") if caption else "") \
                + look_notes(fig.get("state"))
    return n


def _image_slides(deck, pages, n0):
    """One slide per camera sheet / SnapMap site: ``pages`` is
    ``[{"title", "png", "notes"}]`` (pictures sized ``FIGURE_SIZE``)."""
    n = n0
    for pg in pages:
        n += 1
        slide = deck.content_slide(pg["title"], n)
        width, height = FIGURE_SIZE
        pic = slide.shapes.add_picture(
            io.BytesIO(pg["png"]), deck.Inches((SLIDE_W - width) / 2),
            deck.Inches(1.15), width=deck.Inches(width))
        if pic.height > deck.Inches(height):
            ratio = deck.Inches(height) / pic.height
            pic.height = deck.Inches(height)
            pic.width = int(pic.width * ratio)
            pic.left = int((deck.Inches(SLIDE_W) - pic.width) / 2)
        slide.notes_slide.notes_text_frame.text = pg.get("notes", "")
    return n


def build_deck(path, details, logo, file_rows, docs, figures, render_images,
               sections=SECTIONS, image_pages=None, spec=None,
               cover_data=None, notes=None):
    """Write the .pptx to ``path``; returns the number of slides.

    ``spec`` (see ``reportspec``) says which sections go in, in which order,
    which figures and files, and the cover picture; without it ``sections``
    (the old names) do. ``cover_data`` is ``(energy, counts)`` for the "your
    data" cover; a problem with the cover picture is appended to ``notes``.
    ``figures``: ``[{"name", "caption", "state"}]``;
    ``render_images(number, figure)`` returns one PNG (bytes) per page of that
    figure, sized ``FIGURE_SIZE`` inches. ``image_pages()`` returns the camera
    and SnapMap slides as ``[{"title", "png", "notes"}]`` (None: none)."""
    if spec is None:
        spec = reportspec.spec_from_sections(sections, "deck")
    items = reportspec.active(spec)
    sha = reportspec.option(spec, "sha")
    title = (details.get("title") or "").strip() or "Experiment report"
    deck = _Deck(title)
    n = 0
    ids = [sid for sid, _skip in items]
    methods = (details.get("methods") or "").strip()
    cal = (details.get("calibration") or "").strip()
    # the calibration statement rides on the summary slide when both are in
    # (as it always did), else it gets a slide of its own; never twice
    cal_text = cal if cal and not ("methods" in ids and cal in methods) else ""
    for sid, skip in items:
        if sid == "cover":
            art = covers.art(reportspec.cover_of(spec), "pptx", cover_data)
            if art.note and notes is not None:
                notes.append(art.note)
            _title_slide(deck, details, logo, art)
            n += 1
        elif sid == "summary":
            paras = paragraphs(details.get("summary"))
            if cal_text and "calibration" in ids:
                paras.append("Energy calibration: " + cal_text)
            n = _text_slides(deck, "Summary", paras, n, size=16)
        elif sid == "methods":
            n = _text_slides(deck, "Methods", paragraphs(methods), n, size=14)
        elif sid == "calibration":
            if cal_text and "summary" not in ids:
                n = _text_slides(deck, "Energy calibration", [cal_text], n)
        elif sid == "files" and file_rows:
            n = _files_slides(deck, file_rows, n, sha)
        elif sid == "metadata":
            n = _metadata_slides(deck, docs, n, skip)
        elif sid == "images" and image_pages is not None:
            n = _image_slides(deck, image_pages(), n)
        elif sid == "figures" and figures:
            chosen = [f for i, f in enumerate(figures, 1)
                      if reportspec.figure_id(f, i) not in skip]
            n = _figure_slides(deck, chosen, render_images, n)
    if not deck.prs.slides:
        raise PptxError("Nothing to put in the presentation: choose at least "
                        "one section that has content.")
    cp = deck.prs.core_properties
    cp.title = title
    cp.author = (details.get("operator") or "").strip() or appinfo.NAME
    deck.prs.save(path)
    return len(deck.prs.slides)
