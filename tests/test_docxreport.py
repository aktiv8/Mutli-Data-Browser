"""The Word (.docx) export: same data as the PDF report and the PowerPoint
deck, laid out as an editable document.

Run:  python -m unittest discover tests
"""

import io
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import reportspec as rs  # noqa: E402

try:
    import docx  # noqa: F401
    import matplotlib  # noqa: F401
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False


@unittest.skipUnless(HAVE_DOCX, "python-docx / matplotlib not installed")
class TestDocx(unittest.TestCase):
    def setUp(self):
        from test_metasummary import Doc, region
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.docs = [Doc([region("Survey", 160, step=1.0)], "a.vgd"),
                     Doc([region("Mo 3d", 40)], "b.vgd")]
        self.rows = [{"name": "a.vgd", "format": "Thermo", "regions": 1,
                      "size": 2048, "sha256": "ab" * 32}]
        self.details = {"title": "The Study", "summary": "Everything worked.",
                        "methods": "Spectra were recorded.",
                        "calibration": "C 1s at 284.8 eV.",
                        "customer": "Acme", "operator": "J. Bloggs"}

    def figs(self, n):
        return [{"name": f"Fig {i}", "caption": f"Caption {i}", "state": {}}
                for i in range(1, n + 1)]

    def png(self, number, fig):
        from matplotlib.figure import Figure
        f = Figure(figsize=(6.0, 3.0))
        f.text(0.5, 0.5, f"fig {number}")
        buf = io.BytesIO()
        f.savefig(buf, format="png", dpi=30)
        return [buf.getvalue()]

    def build(self, spec=None, figures=3, **kw):
        import docx_export
        path = os.path.join(self.dir, "d.docx")
        self.n = docx_export.build_document(
            path, self.details, "", self.rows, self.docs, self.figs(figures),
            self.png, spec=spec or rs.default_spec(), **kw)
        from docx import Document
        self.doc = Document(path)
        return self.doc

    @staticmethod
    def headings(doc):
        return [p.text for p in doc.paragraphs
                if p.style.name == "Title" or p.style.name.startswith(
                    "Heading")]

    @staticmethod
    def all_text(doc):
        cells = "\n".join(c.text for t in doc.tables for row in t.rows
                          for c in row.cells)
        return "\n".join(p.text for p in doc.paragraphs) + "\n" + cells

    def test_cover_title_and_details(self):
        doc = self.build()
        heads = self.headings(doc)
        self.assertIn("The Study", heads)
        text = self.all_text(doc)
        self.assertIn("Acme", text)
        self.assertIn("J. Bloggs", text)

    def test_contents_has_a_toc_field(self):
        doc = self.build()
        self.assertIn("Contents", self.headings(doc))
        self.assertIn("TOC", doc.element.body.xml)

    def test_summary_methods_and_calibration_sections(self):
        doc = self.build()
        heads = self.headings(doc)
        self.assertIn("Summary", heads)
        self.assertIn("Methods", heads)
        text = self.all_text(doc)
        self.assertIn("Everything worked.", text)
        self.assertIn("Spectra were recorded.", text)
        # the calibration rides on the summary when both are in: no separate
        # "Energy calibration" heading, but the statement is still there
        self.assertNotIn("Energy calibration", heads)
        self.assertIn("C 1s at 284.8 eV.", text)

    def test_calibration_gets_its_own_heading_without_summary(self):
        # the statement rides on the summary when both are in (as the PDF
        # and the deck do); with no summary it gets a heading of its own
        spec = rs.with_on(rs.default_spec(), "summary", False)
        doc = self.build(spec)
        self.assertIn("Energy calibration", self.headings(doc))
        self.assertIn("C 1s at 284.8 eV.", self.all_text(doc))

    def test_files_table(self):
        doc = self.build()
        self.assertIn("Data files", self.headings(doc))
        rows = [[c.text for c in row.cells] for t in doc.tables
               for row in t.rows]
        self.assertTrue(any("a.vgd" in r for r in rows))

    def test_figures_are_embedded_with_captions(self):
        spec = rs.with_on(rs.default_spec(), "cover", False)
        doc = self.build(spec, figures=3)
        heads = self.headings(doc)
        self.assertIn("Figure 1 — Fig 1", heads)
        self.assertIn("Figure 3 — Fig 3", heads)
        self.assertEqual(len(doc.inline_shapes), 3)
        self.assertIn("Caption 1", self.all_text(doc))

    def test_a_section_can_be_switched_off(self):
        spec = rs.with_on(rs.default_spec(), "figures", False)
        spec = rs.with_on(spec, "cover", False)
        doc = self.build(spec, figures=2)
        self.assertNotIn("Figures", self.headings(doc))
        self.assertEqual(len(doc.inline_shapes), 0)

    def test_nothing_selected_raises(self):
        import docx_export
        spec = rs.with_all(rs.default_spec(), False)
        path = os.path.join(self.dir, "empty.docx")
        with self.assertRaises(docx_export.DocxError):
            docx_export.build_document(
                path, self.details, "", self.rows, self.docs, [], self.png,
                spec=spec)

    def test_footer_has_page_fields(self):
        doc = self.build()
        footer = doc.sections[0].footer
        xml = footer.part.element.xml
        self.assertIn("PAGE", xml)
        self.assertIn("NUMPAGES", xml)


if __name__ == "__main__":
    unittest.main()
