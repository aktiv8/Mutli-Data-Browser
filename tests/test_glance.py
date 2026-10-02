"""The "At a glance" and "Timing" report sections: their facts (glance.py),
the spec rules that govern them (reportspec.py: old reports do not gain
them, a tick removes them) and how the PDF, the slides and the Word document
lay them out.

Run:  python -m unittest discover tests
"""

import os
import shutil
import sys
import tempfile
import unittest
from types import SimpleNamespace

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import glance  # noqa: E402
import reportspec as rs  # noqa: E402
import timing  # noqa: E402
from test_timing import doc, region  # noqa: E402

try:
    import pymupdf as mupdf
except ImportError:
    try:
        import fitz as mupdf
    except ImportError:
        mupdf = None
try:
    import reportlab  # noqa: F401
    import matplotlib  # noqa: F401
    HAVE_PDF = mupdf is not None
except ImportError:
    HAVE_PDF = False
try:
    import pptx  # noqa: F401
    HAVE_PPTX = HAVE_PDF
except ImportError:
    HAVE_PPTX = False
try:
    import docx  # noqa: F401
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

MD = [{"Sample": "Film A", "Instrument": "Thermo K-Alpha+",
       "Operator": "D. Morgan", "Date acquired": "2026-08-24"},
      {"Sample": "Film B", "Instrument": "Thermo K-Alpha+",
       "Operator": "D. Morgan", "Date acquired": "2026-08-26"},
      {"Sample": "Film B", "Instrument": "Thermo K-Alpha+",
       "Operator": "D. Morgan", "Date acquired": "2026-08-26"}]
FILES = [{"name": "a.vgd", "format": "Thermo Avantage"},
         {"name": "b.vgd", "format": "Thermo Avantage"}]


def runs(start="2026-08-24 09:00:00", end="2026-08-24 10:00:00", **kw):
    return region(start=start, end=end, **kw)


def fake_results(*samples):
    """A ``resultspages.Results``-shaped object: each sample is a list of
    (region name, at %) for one level."""
    out = []
    for regs in samples:
        entries = [{"spectrum": n, "row": {"region": n}} for n, _p in regs]
        res = [{"at_pct": p} for _n, p in regs]
        out.append(SimpleNamespace(levels=[SimpleNamespace(entries=entries,
                                                           res=res)],
                                   casaxps=None))
    return SimpleNamespace(samples=out)


class TestFacts(unittest.TestCase):
    def summary(self):
        return timing.summarise([doc(runs(), runs(
            start="2026-08-24 12:00:00", end="2026-08-24 13:00:00"))])

    def test_everything_recorded_gives_every_line(self):
        got = dict(glance.facts(
            MD, FILES, fake_results([("C 1s", 60.0), ("O 1s", 40.0)],
                                    [("Ti 2p", 100.0)]),
            self.summary(), figures=4, pictures=11))
        self.assertEqual(got["Data"], "2 files · 2 samples · 3 spectra")
        self.assertEqual(got["Format"], "Thermo Avantage")
        self.assertEqual(got["Instrument"], "Thermo K-Alpha+")
        self.assertEqual(got["Operator"], "D. Morgan")
        self.assertEqual(got["Acquired"], "2026-08-24 – 2026-08-26")
        self.assertEqual(got["Counting time"],
                         "4 min 11 s  (instrument in use 2 h 00 min)")
        self.assertEqual(got["Quantified"], "C, O, Ti  (2 samples)")
        self.assertEqual(got["Also included"],
                         "4 saved figures · 11 camera pictures and maps")
        self.assertEqual(list(got), ["Data", "Format", "Instrument",
                                     "Operator", "Acquired", "Counting time",
                                     "Quantified", "Also included"])

    def test_a_line_is_only_written_when_it_is_recorded(self):
        md = [{"Sample": "S", "Instrument": "(unknown)"}]
        got = dict(glance.facts(md, [], None, None))
        self.assertEqual(got, {"Data": "1 sample · 1 spectrum"})
        self.assertEqual(glance.facts([], [], None, None), [])
        self.assertEqual(glance.facts(None, None), [])

    def test_one_of_each_is_singular(self):
        got = dict(glance.facts([{"Sample": "S"}], [FILES[0]], None, None,
                                figures=1, pictures=1))
        self.assertEqual(got["Data"],
                         "1 file · 1 sample · 1 spectrum")
        self.assertEqual(got["Also included"],
                         "1 saved figure · 1 camera picture or map")

    def test_counting_time_needs_every_scan_count(self):
        known = self.summary()
        self.assertIn("instrument in use", glance.counting_text(known))
        missing = timing.summarise([doc(runs(), runs(scans=None))])
        self.assertEqual(glance.counting_text(missing), "")
        self.assertEqual(glance.counting_text(None), "")
        # starts only (a Kratos file): the counting time, no "in use"
        a = region(start="2026-09-12 09:48:53", tz="")
        only = timing.summarise([doc(a)])
        self.assertEqual(glance.counting_text(only), "2 min 06 s")

    def test_unquantified_regions_do_not_name_an_element(self):
        results = fake_results([("C 1s", 60.0), ("O 1s", None)])
        self.assertEqual(glance.quantified_elements(results),
                         (["C"], 1))
        self.assertEqual(glance.quantified_elements(None), ([], 0))

    def test_casaxps_tables_name_their_elements(self):
        cq = SimpleNamespace(survey=[{"element": "O 1s", "pct": 1.8},
                                     {"element": "C 1s", "pct": 27.4}])
        results = SimpleNamespace(samples=[SimpleNamespace(levels=[],
                                                           casaxps=cq)])
        self.assertEqual(glance.quantified_elements(results), (["O", "C"], 1))


class TestTimingSection(unittest.TestCase):
    def test_rows_notes_and_the_per_sample_table(self):
        a = runs(sample="Film A")
        b = runs(sample="Film B", scans=20)
        d1 = doc(a, b)
        d1.path = "exp1.vgd"
        sec = glance.timing_section([d1])
        self.assertTrue(sec)
        labels = [k for k, _v in sec.rows]
        self.assertEqual(labels, ["First start", "Last finish",
                                  "First start to last finish",
                                  "Instrument in use", "Counting time",
                                  "Not counting (moves, settling, "
                                  "sputtering, dead time)"])
        self.assertEqual(sec.by_sample, [("Film A", "exp1.vgd", "2 min 06 s"),
                                         ("Film B", "exp1.vgd", "4 min 11 s")])
        self.assertEqual(sec.notes, [])

    def test_one_sample_needs_no_table(self):
        sec = glance.timing_section([doc(runs(sample="Only"))])
        self.assertEqual(sec.by_sample, [])
        self.assertTrue(sec.rows)

    def test_starts_only_gives_the_short_version_and_says_why(self):
        a = region(start="2026-09-12 09:48:53", tz="")
        b = region(start="2026-09-13 16:41:00", tz="")
        sec = glance.timing_section([doc(a, b)])
        self.assertEqual([k for k, _v in sec.rows],
                         ["First start", "Last start", "Counting time"])
        self.assertTrue(any("started" in n for n in sec.notes))

    def test_no_times_at_all_is_empty(self):
        sec = glance.timing_section([doc(region(scans=None))])
        self.assertFalse(sec)
        self.assertTrue(sec.notes)
        self.assertFalse(glance.timing_section([]))

    def test_a_sample_with_no_scan_count_says_not_recorded(self):
        d1 = doc(runs(sample="A"), runs(sample="B", scans=None))
        sec = glance.timing_section([d1])
        self.assertIn(("B", "x.vgd", "not recorded"), sec.by_sample)

    def test_the_sample_names_follow_the_users_renames(self):
        d1 = doc(runs(sample="raw1"), runs(sample="raw2"))
        d1.file_id = "f1"
        d1.annotations = SimpleNamespace(
            sample_label=lambda fid, s: {"raw1": "Film A"}.get(s, s))
        names = [s for s, _f, _t in glance.timing_section([d1]).by_sample]
        self.assertEqual(names, ["Film A", "raw2"])


class TestSpecRules(unittest.TestCase):
    def test_they_sit_where_a_reader_looks_for_them(self):
        order = rs.order(rs.default_spec())
        self.assertEqual(order[:4], ["cover", "contents", "glance", "summary"])
        self.assertEqual(order.index("timing"), order.index("methods") + 1)
        self.assertTrue(rs.is_on(rs.default_spec(), "glance"))
        self.assertTrue(rs.is_on(rs.default_spec(), "timing"))

    def test_a_report_saved_before_them_does_not_gain_pages(self):
        old = {"sections": [{"id": i, "on": True} for i in (
            "cover", "contents", "summary", "results", "figures", "images",
            "methods", "calibration", "metadata", "files")]}
        s = rs.sanitise(old)
        self.assertFalse(rs.is_on(s, "glance"))
        self.assertFalse(rs.is_on(s, "timing"))
        self.assertEqual(rs.order(s)[:4],
                         ["cover", "contents", "glance", "summary"])
        # but a spec the user switches on keeps it on, and re-sanitising is
        # stable (it is not migrated twice)
        s = rs.with_on(s, "timing", True)
        self.assertTrue(rs.is_on(rs.sanitise(s), "timing"))
        self.assertFalse(rs.is_on(rs.sanitise(s), "glance"))

    def test_garbage_and_empty_input_give_the_plain_default(self):
        for bad in (None, 5, {"sections": 3}, {"sections": []}):
            self.assertTrue(rs.is_on(rs.sanitise(bad), "glance"), bad)
            self.assertTrue(rs.is_on(rs.sanitise(bad), "timing"), bad)

    def test_the_old_section_names_never_switch_them_on(self):
        for kind in ("pdf", "deck"):
            spec = rs.spec_from_sections(
                ("cover", "title", "files", "metadata", "images", "figures"),
                kind)
            self.assertFalse(rs.is_on(spec, "glance"))
            self.assertFalse(rs.is_on(spec, "timing"))

    def test_presets(self):
        on = lambda name, sid: rs.is_on(rs.BUILTIN_PRESETS[name], sid)  # noqa: E731
        self.assertTrue(on("Customer report", "glance"))
        self.assertFalse(on("Customer report", "timing"))
        self.assertTrue(on("Audit trail", "timing"))
        self.assertFalse(on("Audit trail", "glance"))
        for sid in ("glance", "timing"):
            self.assertFalse(on("Quick look", sid))
            self.assertTrue(on("Everything", sid))

    def test_the_inventory_says_what_is_missing(self):
        args = dict(details={}, methods_text="", calibration="",
                    file_rows=[], docs=[], figures=[], has_images=False)
        inv = rs.inventory(**args)
        self.assertNotIn("glance", inv.present_ids())
        self.assertNotIn("timing", inv.present_ids())
        self.assertIn("acquisition times", inv.summary("timing"))
        inv = rs.inventory(glance=[("Data", "1 file")],
                           timing=glance.Timing(rows=[("a", "b")]), **args)
        self.assertLessEqual({"glance", "timing"}, inv.present_ids())
        # a caller that knows nothing of the new sections still gets
        # "At a glance" when there are files, and no Timing
        inv = rs.inventory(**dict(args, file_rows=[{"name": "a"}]))
        self.assertIn("glance", inv.present_ids())
        self.assertNotIn("timing", inv.present_ids())

    def test_describe_names_them(self):
        inv = rs.inventory(details={"summary": "s"}, methods_text="m",
                           calibration="", file_rows=[{"name": "a"}],
                           docs=[], figures=[], has_images=False,
                           glance=[("Data", "1 file")],
                           timing=glance.Timing(rows=[("a", "b")]))
        text = rs.describe(rs.default_spec(), inv)
        self.assertIn("At a glance", text)
        self.assertIn("Timing", text)
        off = rs.with_on(rs.with_on(rs.default_spec(), "glance", False),
                         "timing", False)
        text = rs.describe(off, inv)
        self.assertNotIn("At a glance", text)
        self.assertNotIn("Timing", text)


def sample_facts():
    return glance.facts(MD, FILES, None, timing.summarise([doc(
        runs(sample="Film A"), runs(sample="Film B", scans=20))]),
        figures=2)


def sample_timing():
    d1 = doc(runs(sample="Film A"), runs(sample="Film B", scans=20))
    d1.path = "exp1.vgd"
    return glance.timing_section([d1])


class Built(unittest.TestCase):
    def setUp(self):
        from test_metasummary import Doc, region as mregion
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.docs = [Doc([mregion("Survey", 160, step=1.0)], "a.vgd")]
        self.rows = [{"name": "a.vgd", "format": "Thermo", "regions": 1,
                      "size": 2048, "sha256": "ab" * 32}]
        self.details = {"title": "Study", "summary": "Everything worked.",
                        "methods": "Spectra were recorded.",
                        "calibration": "C 1s at 284.8 eV."}
        self.facts, self.timing = sample_facts(), sample_timing()


@unittest.skipUnless(HAVE_PDF, "reportlab / PyMuPDF / matplotlib missing")
class TestPdf(Built):
    def build(self, spec=None, **kw):
        import report
        self.k = getattr(self, "k", 0) + 1       # Windows: one file per build
        path = os.path.join(self.dir, f"r{self.k}.pdf")
        kw.setdefault("glance", self.facts)
        kw.setdefault("timing", self.timing)
        report.build_report(path, self.details, "", self.rows, self.docs, [],
                            lambda *a: 0, spec=spec or rs.default_spec(), **kw)
        self.doc = mupdf.open(path)
        self.addCleanup(self.doc.close)
        return [p.get_text() for p in self.doc]

    def test_both_sections_are_there_with_their_lines(self):
        text = "\n".join(self.build())
        for token in ("At a glance", "Thermo K-Alpha+", "D. Morgan",
                      "2026-08-24", "Timing", "First start",
                      "Instrument in use", "Counting time by sample",
                      "Film A", "exp1.vgd"):
            self.assertIn(token, text)

    def test_they_are_in_the_contents_and_the_bookmarks_on_the_right_page(self):
        pages = self.build()
        toc = {t: p for _l, t, p in self.doc.get_toc()}
        self.assertIn("At a glance", toc)
        self.assertIn("Timing", toc)
        self.assertLess(toc["At a glance"], toc["Timing"])
        self.assertIn("At a glance", pages[toc["At a glance"] - 1])
        self.assertIn("First start", pages[toc["Timing"] - 1])
        self.assertIn("At a glance", pages[1])             # printed contents
        self.assertIn("Timing", pages[1])
        links = {self.doc[1].get_textbox(l["from"]).strip(): l["page"] + 1
                 for l in self.doc[1].get_links()}
        self.assertEqual(links["At a glance"], toc["At a glance"])

    def test_a_tick_removes_each_section(self):
        for sid, token in (("glance", "Thermo K-Alpha+"),
                           ("timing", "Instrument in use")):
            pages = self.build(rs.with_on(rs.default_spec(), sid, False))
            text = "\n".join(pages)
            self.assertNotIn(token, text, sid)
            self.assertNotIn(rs.LABELS[sid], pages[1], sid)  # nor in contents

    def test_they_come_in_the_order_of_the_spec(self):
        spec = rs.moved(rs.default_spec(), "timing", -3)
        order = rs.order(spec)
        self.assertLess(order.index("timing"), order.index("methods"))
        self.build(spec)
        names = [t for _l, t, _p in self.doc.get_toc()]
        self.assertLess(names.index("Timing"), names.index("Methods"))

    def test_nothing_to_show_leaves_the_section_out(self):
        pages = self.build(glance=[], timing=glance.Timing())
        text = "\n".join(pages)
        self.assertNotIn("At a glance", text)
        self.assertNotIn("Counting time by sample", text)
        pages = self.build(glance=None, timing=None)
        self.assertNotIn("At a glance", "\n".join(pages))

    def test_starts_only_timing_prints_its_caveat(self):
        a = region(start="2026-09-12 09:48:53", tz="")
        sec = glance.timing_section([doc(a)])
        sec.notes.append("Only when each run started is recorded.")
        text = "\n".join(self.build(timing=sec))
        self.assertIn("Only when each run started is recorded.", text)
        self.assertNotIn("Last finish", text)


@unittest.skipUnless(HAVE_PPTX, "python-pptx / PyMuPDF not installed")
class TestDeck(Built):
    def build(self, spec=None, **kw):
        import pptx_export
        from pptx import Presentation
        self.k = getattr(self, "k", 0) + 1
        path = os.path.join(self.dir, f"d{self.k}.pptx")
        kw.setdefault("glance", self.facts)
        kw.setdefault("timing", self.timing)
        pptx_export.build_deck(path, self.details, "", self.rows, self.docs,
                               [], lambda n, f: [],
                               spec=spec or rs.default_spec(), **kw)
        return list(Presentation(path).slides)

    @staticmethod
    def title(slide):
        return slide.shapes.title.text if slide.shapes.title is not None \
            else ""

    @staticmethod
    def table_cells(slide):
        return [[c.text for c in row.cells] for sh in slide.shapes
                if sh.has_table for row in sh.table.rows]

    def test_a_slide_each_with_the_same_lines(self):
        slides = self.build()
        titles = [self.title(s) for s in slides]
        glance_slide = slides[titles.index("At a glance")]
        cells = dict(tuple(r) for r in self.table_cells(glance_slide))
        self.assertEqual(cells["Operator"], "D. Morgan")
        timing_slide = slides[titles.index("Timing")]
        rows = dict(tuple(r) for r in self.table_cells(timing_slide))
        self.assertIn("Instrument in use", rows)
        by = slides[titles.index("Timing: counting time by sample")]
        table = self.table_cells(by)
        self.assertEqual(table[0], ["Sample", "File", "Counting time"])
        self.assertEqual(table[1][:2], ["Film A", "exp1.vgd"])
        self.assertLess(titles.index("At a glance"), titles.index("Summary"))
        self.assertLess(titles.index("Methods"), titles.index("Timing"))

    def test_they_are_in_the_contents_with_true_numbers(self):
        slides = self.build()
        contents = next(s for s in slides if self.title(s) == "Contents")
        rows = {r[0].strip(): int(r[1]) for r in self.table_cells(contents)
                if r[1].strip().isdigit()}
        for name in ("At a glance", "Timing"):
            self.assertIn(name, rows)
            self.assertEqual(self.title(slides[rows[name] - 1]), name)

    def test_a_tick_removes_each_section(self):
        for sid in ("glance", "timing"):
            titles = [self.title(s) for s in self.build(
                rs.with_on(rs.default_spec(), sid, False))]
            self.assertNotIn(rs.LABELS[sid], titles, sid)
            self.assertFalse([t for t in titles if t.startswith(
                rs.LABELS[sid] + ":")], sid)

    def test_a_note_goes_under_the_timing_table(self):
        sec = glance.timing_section([doc(region(start="2026-09-12 09:48:53",
                                                tz=""))])
        slides = self.build(timing=sec)
        slide = next(s for s in slides if self.title(s) == "Timing")
        text = " ".join(sh.text_frame.text for sh in slide.shapes
                        if sh.has_text_frame)
        self.assertIn("started", text)
        self.assertFalse([s for s in slides
                          if self.title(s).startswith("Timing:")])


@unittest.skipUnless(HAVE_DOCX, "python-docx not installed")
class TestDocx(Built):
    def build(self, spec=None, **kw):
        import docx_export
        from docx import Document
        self.k = getattr(self, "k", 0) + 1
        path = os.path.join(self.dir, f"d{self.k}.docx")
        kw.setdefault("glance", self.facts)
        kw.setdefault("timing", self.timing)
        docx_export.build_document(
            path, self.details, "", self.rows, self.docs, [],
            lambda n, f: [], spec=spec or rs.default_spec(), **kw)
        return Document(path)

    @staticmethod
    def headings(doc):
        return [p.text for p in doc.paragraphs
                if p.style.name == "Title" or p.style.name.startswith(
                    "Heading")]

    @staticmethod
    def all_text(doc):
        cells = "\n".join(c.text for t in doc.tables for row in t.rows
                          for c in row.cells)
        return "\n".join(p.text for p in doc.paragraphs) + "\n" + cells

    def test_both_sections_are_headings_with_their_tables(self):
        doc = self.build()
        heads = self.headings(doc)
        self.assertLess(heads.index("At a glance"), heads.index("Summary"))
        self.assertLess(heads.index("Methods"), heads.index("Timing"))
        text = self.all_text(doc)
        for token in ("D. Morgan", "Instrument in use",
                      "Counting time by sample", "exp1.vgd"):
            self.assertIn(token, text)

    def test_a_tick_removes_each_section(self):
        for sid in ("glance", "timing"):
            doc = self.build(rs.with_on(rs.default_spec(), sid, False))
            self.assertNotIn(rs.LABELS[sid], self.headings(doc), sid)

    def test_nothing_to_show_leaves_them_out(self):
        doc = self.build(glance=None, timing=None)
        heads = self.headings(doc)
        self.assertNotIn("At a glance", heads)
        self.assertNotIn("Timing", heads)


if __name__ == "__main__":
    unittest.main()
