"""Cover pictures for the report and the slides (milestone G3).

Run:  python -m unittest discover tests
"""

import io
import os
import shutil
import struct
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import appinfo  # noqa: E402
import covers  # noqa: E402
import reportspec as rs  # noqa: E402

try:
    from PIL import Image
    HAVE_PIL = True
except ImportError:
    HAVE_PIL = False


def png_size(png):
    return struct.unpack(">II", png[16:24])


class Tmp(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)

    def picture(self, name, size=(300, 100), colour=(200, 40, 40), fmt="PNG"):
        path = os.path.join(self.dir, name)
        Image.new("RGB", size, colour).save(path, format=fmt)
        return path


class TestColour(unittest.TestCase):
    def test_accents(self):
        self.assertEqual(covers.valid_accent("#1f7a8c"), "#1F7A8C")
        for bad in ("", None, "1f7a8c", "#12345", "#GGGGGG", "red", 5):
            self.assertEqual(covers.valid_accent(bad), "")

    def test_tints_run_from_the_colour_to_white(self):
        self.assertEqual(covers.tint("#2C3E50", 0), "#2C3E50")
        self.assertEqual(covers.tint("#2C3E50", 1), "#FFFFFF")
        self.assertEqual(covers.mix("#000000", "#FFFFFF", 0.5), "#808080")


class TestList(Tmp):
    def test_the_list_starts_with_none_and_the_built_ins(self):
        ids = [c.id for c in covers.list_covers(self.dir)]
        self.assertEqual(ids[:6], ["none", "ribbon", "band", "minimal", "data",
                                   "grid"])

    @unittest.skipUnless(HAVE_PIL, "Pillow not installed")
    def test_pictures_in_the_folder_are_offered_by_name(self):
        self.picture("Lab logo.png")
        self.picture("notes.txt.jpg", fmt="JPEG")
        open(os.path.join(self.dir, "readme.txt"), "w").close()
        got = [(c.id, c.name) for c in covers.list_covers(self.dir)
               if c.kind == "file"]
        self.assertEqual(got, [("file:Lab logo.png", "Lab logo"),
                               ("file:notes.txt.jpg", "notes.txt")])

    def test_the_splash_picture_is_not_offered(self):
        # it does not crop well to a strip: a picture of your own does
        self.assertNotIn("splash", [c.id for c in covers.list_covers(self.dir)])
        self.assertEqual(covers.resolve({"design": "splash"})[0], "none")

    def test_a_missing_folder_is_an_empty_list(self):
        self.assertEqual(appinfo.cover_images(os.path.join(self.dir, "nope")),
                         [])
        self.assertTrue(appinfo.cover_dir().endswith(
            os.path.join("assets", "covers")))


class TestResolve(Tmp):
    def test_designs_and_their_notes(self):
        self.assertEqual(covers.resolve({"design": "band"}, self.dir),
                         ("band", None, ""))
        self.assertEqual(covers.resolve({}, self.dir)[0], "none")
        self.assertEqual(covers.resolve({"design": "none"}, self.dir)[0],
                         "none")
        d, p, note = covers.resolve({"design": "wat"}, self.dir)
        self.assertEqual((d, p), ("none", None))
        self.assertIn("unknown", note)

    def test_a_picture_that_is_gone_says_so(self):
        d, _p, note = covers.resolve({"design": "file:gone.png"}, self.dir)
        self.assertEqual(d, "none")
        self.assertIn("gone.png", note)
        d, _p, note = covers.resolve({"design": "image", "image": "/x/y.png"},
                                     self.dir)
        self.assertEqual(d, "none")
        self.assertIn("y.png", note)

    @unittest.skipUnless(HAVE_PIL, "Pillow not installed")
    def test_a_picture_in_the_folder_and_a_browsed_one(self):
        p = self.picture("a.png")
        self.assertEqual(covers.resolve({"design": "file:a.png"}, self.dir),
                         ("picture", p, ""))
        self.assertEqual(covers.resolve({"design": "image", "image": p}),
                         ("picture", p, ""))


@unittest.skipUnless(covers.HAVE_MPL, "matplotlib not installed")
class TestBuiltInDesigns(unittest.TestCase):
    DATA = ([300.0 - 0.5 * i for i in range(200)],
            [100 + 900 * (1 if 95 < i < 105 else 0) + i for i in range(200)])

    def test_every_design_fits_its_space_for_both_outputs(self):
        for design, _n in covers.BUILTIN:
            for kind, (w, h, unit, dpi) in covers.SIZES.items():
                a = covers.art({"design": design}, kind, data=self.DATA)
                inches = (w / 25.4, h / 25.4) if unit == "mm" else (w, h)
                want = (round(inches[0] * dpi), round(inches[1] * dpi))
                self.assertIsNotNone(a.data, (design, kind))
                self.assertEqual(png_size(a.data), want, (design, kind))
                self.assertEqual((a.width, a.height, a.unit), (w, h, unit))
                self.assertEqual(a.note, "")

    @unittest.skipUnless(HAVE_PIL, "Pillow not installed")
    def test_designs_are_not_blank_and_differ_from_each_other(self):
        seen = set()
        for design, _n in covers.BUILTIN:
            png = covers.art({"design": design}, "pdf", data=self.DATA).data
            im = Image.open(io.BytesIO(png)).convert("L")
            lo, hi = im.getextrema()
            self.assertGreater(hi - lo, 40, design)          # something drawn
            seen.add(png)
        self.assertEqual(len(seen), len(covers.BUILTIN))

    def test_the_same_cover_is_drawn_the_same_every_time(self):
        covers._cache.clear()
        a = covers.art({"design": "ribbon"}, "pdf").data
        covers._cache.clear()
        b = covers.art({"design": "ribbon"}, "pdf").data
        self.assertEqual(a, b)
        covers._cache.clear()
        g1 = covers.art({"design": "grid"}, "pdf").data
        covers._cache.clear()
        self.assertEqual(g1, covers.art({"design": "grid"}, "pdf").data)

    def test_the_accent_changes_the_picture(self):
        navy = covers.art({"design": "band"}, "pdf").data
        teal = covers.art({"design": "band", "accent": "#1F7A8C"}, "pdf").data
        self.assertNotEqual(navy, teal)
        # a bad accent means the default, not an error
        self.assertEqual(
            covers.art({"design": "band", "accent": "nonsense"}, "pdf").data,
            navy)

    def test_the_data_design_draws_the_data_and_falls_back_without_it(self):
        with_data = covers.art({"design": "data"}, "pdf", data=self.DATA).data
        other = ([0, 1, 2, 3], [5, 1, 9, 2])
        self.assertNotEqual(with_data, covers.art({"design": "data"}, "pdf",
                                                  data=other).data)
        ribbon = covers.art({"design": "ribbon"}, "pdf").data
        for bad in (None, ([], []), ([1, 2], [1, 2]), ("x", 3),
                    ([1, 2, 3], [1, float("nan"), 3])):
            self.assertEqual(covers.art({"design": "data"}, "pdf",
                                        data=bad).data, ribbon, bad)

    def test_none_is_no_picture(self):
        a = covers.art({"design": "none"}, "pdf")
        self.assertIsNone(a.data)
        self.assertEqual(a.note, "")

    def test_a_thumbnail_is_a_small_png(self):
        png = covers.thumbnail({"design": "ribbon"})
        self.assertEqual(png_size(png), (240, 60))
        self.assertIsNone(covers.thumbnail({"design": "none"}))


@unittest.skipUnless(HAVE_PIL, "Pillow not installed")
class TestPictures(Tmp):
    def test_a_wide_picture_is_cropped_from_the_middle_not_stretched(self):
        path = os.path.join(self.dir, "wide.png")
        im = Image.new("RGB", (400, 100), (255, 255, 255))
        for x in range(150, 250):                    # a red block in the middle
            for y in range(100):
                im.putpixel((x, y), (255, 0, 0))
        im.save(path)
        art = covers.art({"design": "image", "image": path}, "pdf")
        self.assertEqual(art.fmt, "jpeg")
        self.assertEqual(art.data[:2], b"\xff\xd8")
        out = Image.open(io.BytesIO(art.data)).convert("RGB")
        self.assertEqual(out.size, (1417, 354))
        red = out.getpixel((out.width // 2, out.height // 2))
        self.assertGreater(red[0], 235)                      # the middle survived
        self.assertLess(red[1] + red[2], 40)
        self.assertGreater(min(out.getpixel((5, 5))), 235)   # white at the edge

    def test_a_tall_picture_and_a_tiny_one_are_scaled_up_to_fill(self):
        for size in ((100, 400), (3, 2)):
            path = self.picture("p.png", size=size)
            art = covers.art({"design": "image", "image": path}, "pptx")
            self.assertEqual(Image.open(io.BytesIO(art.data)).size,
                             (2000, 232))

    def test_a_photograph_is_small_and_the_preview_is_a_png(self):
        path = os.path.join(self.dir, "photo.png")
        import random
        rnd = random.Random(3)
        im = Image.new("RGB", (900, 300))
        im.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256))
                    for _ in range(900 * 300)])
        im.save(path)
        cover = {"design": "image", "image": path}
        art = covers.art(cover, "pdf")
        self.assertEqual(art.fmt, "jpeg")
        self.assertLess(len(art.data), 800_000)             # noise: worst case
        thumb = covers.thumbnail(cover)
        self.assertEqual(thumb[:8], b"\x89PNG\r\n\x1a\n")   # Tk can show it
        self.assertEqual(png_size(thumb), (240, 60))
        self.assertEqual(covers.art({"design": "ribbon"}, "pdf").fmt, "png")

    def test_a_folder_picture_and_the_cache(self):
        path = self.picture("a.jpg", fmt="JPEG", colour=(10, 120, 200))
        cover = {"design": "file:a.jpg"}
        a = covers.art(cover, "pdf", folder=self.dir)
        self.assertIsNotNone(a.data)
        self.assertEqual(covers.art(cover, "pdf", folder=self.dir).data, a.data)
        # replaced on disk: the cover follows
        Image.new("RGB", (300, 100), (200, 10, 10)).save(path, format="JPEG")
        os.utime(path, (1, 1))                       # a different mtime
        self.assertNotEqual(covers.art(cover, "pdf", folder=self.dir).data,
                            a.data)

    def test_an_unreadable_picture_is_a_note_not_an_error(self):
        path = os.path.join(self.dir, "bad.png")
        with open(path, "wb") as fh:
            fh.write(b"this is not a picture")
        a = covers.art({"design": "image", "image": path}, "pdf")
        self.assertIsNone(a.data)
        self.assertIn("bad.png", a.note)
        a = covers.art({"design": "image", "image": "/no/such.png"}, "pdf")
        self.assertIsNone(a.data)
        self.assertIn("not found", a.note)

    def test_without_pillow_a_picture_is_text_only(self):
        path = self.picture("a.png")
        real = covers.HAVE_PIL
        covers.HAVE_PIL = False
        self.addCleanup(setattr, covers, "HAVE_PIL", real)
        covers._cache.clear()
        a = covers.art({"design": "image", "image": path}, "pdf")
        self.assertIsNone(a.data)
        self.assertTrue(a.note)


class TestWithoutMatplotlib(unittest.TestCase):
    def test_built_ins_are_text_only(self):
        real = covers.HAVE_MPL
        covers.HAVE_MPL = False
        self.addCleanup(setattr, covers, "HAVE_MPL", real)
        covers._cache.clear()
        a = covers.art({"design": "ribbon"}, "pdf")
        self.assertIsNone(a.data)
        self.assertIn("matplotlib", a.note)


class TestSpecCover(unittest.TestCase):
    def test_the_default_has_a_cover_and_the_old_sections_have_none(self):
        self.assertEqual(rs.cover_of(rs.default_spec()),
                         {"design": "ribbon", "image": "", "accent": ""})
        self.assertEqual(rs.cover_of(rs.spec_from_sections(("cover",)))
                         ["design"], "none")

    def test_the_cover_is_sanitised(self):
        s = rs.sanitise({"cover": {"design": 5, "image": ["x"],
                                   "accent": "teal"}})
        self.assertEqual(rs.cover_of(s), rs.DEFAULT_COVER)
        s = rs.sanitise({"cover": {"design": "file:a.png", "image": "p.png",
                                   "accent": "#1f7a8c"}})
        self.assertEqual(rs.cover_of(s), {"design": "file:a.png",
                                          "image": "p.png",
                                          "accent": "#1F7A8C"})
        self.assertEqual(rs.cover_of(rs.sanitise({"cover": "x"})),
                         rs.DEFAULT_COVER)
        self.assertEqual(rs.cover_of(rs.sanitise({})), rs.DEFAULT_COVER)

    def test_changing_the_cover_and_a_preset_leaves_it_alone(self):
        s = rs.with_cover(rs.default_spec(), design="band", accent="#1F7A8C")
        self.assertEqual(rs.cover_of(s)["design"], "band")
        quick = rs.with_cover_of(rs.BUILTIN_PRESETS["Quick look"], s)
        self.assertEqual(rs.cover_of(quick), rs.cover_of(s))
        self.assertEqual([i for i, _k in rs.active(quick)],
                         ["cover", "figures"])              # its sections

    def test_an_old_spec_without_a_cover_gets_the_default(self):
        old = {"version": 1, "sections": [{"id": "cover", "on": True}],
               "skip": {}, "options": {"sha": "short"}}
        self.assertEqual(rs.cover_of(old), rs.DEFAULT_COVER)


try:
    import reportlab  # noqa: F401
    try:
        import pymupdf as mupdf
    except ImportError:
        import fitz as mupdf
    HAVE_PDF = True
except ImportError:
    HAVE_PDF = False
try:
    import pptx  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


@unittest.skipUnless(HAVE_PDF and covers.HAVE_MPL,
                     "reportlab / PyMuPDF / matplotlib not installed")
class TestPdfCover(Tmp):
    def build(self, cover, notes=None):
        import report
        spec = rs.with_all(rs.default_spec(), False)
        spec = rs.with_cover(rs.with_on(spec, "cover", True), **cover)
        path = os.path.join(self.dir, "r.pdf")
        report.build_report(path, {"title": "T", "customer": "ACME"}, "", [],
                            [], [], None, spec=spec,
                            cover_data=TestBuiltInDesigns.DATA, notes=notes)
        with mupdf.open(path) as d:
            page = d[0]
            return {"images": page.get_image_info(),
                    "customer": page.search_for("ACME")}

    def test_the_picture_is_on_the_first_page_above_the_title(self):
        got = self.build({"design": "ribbon"})
        self.assertEqual(len(got["images"]), 1)
        top = got["images"][0]["bbox"]
        self.assertLess(top[3], got["customer"][0].y0)    # picture above text
        width_mm = (top[2] - top[0]) / 72 * 25.4
        self.assertAlmostEqual(width_mm, 180.0, delta=1.0)

    def test_no_picture_without_a_design_and_a_note_when_it_is_missing(self):
        self.assertEqual(self.build({"design": "none"})["images"], [])
        notes = []
        got = self.build({"design": "file:gone.png"}, notes)
        self.assertEqual(got["images"], [])                # still a report
        self.assertTrue(any("gone.png" in n for n in notes))


@unittest.skipUnless(HAVE_PPTX and covers.HAVE_MPL,
                     "python-pptx / matplotlib not installed")
class TestDeckCover(Tmp):
    def build(self, cover, notes=None):
        import pptx_export
        from pptx import Presentation
        spec = rs.with_cover(rs.with_on(rs.with_all(rs.default_spec(), False),
                                        "cover", True), **cover)
        path = os.path.join(self.dir, "d.pptx")
        pptx_export.build_deck(path, {"title": "T"}, "", [], [], [], None,
                               spec=spec, notes=notes)
        return Presentation(path).slides[0]

    def pictures(self, slide):
        return [s for s in slide.shapes if s.shape_type == 13]

    def test_the_picture_runs_along_the_bottom_of_the_title_slide(self):
        import pptx_export
        slide = self.build({"design": "band", "accent": "#1F7A8C"})
        pics = self.pictures(slide)
        self.assertEqual(len(pics), 1)
        pic = pics[0]
        emu = 914400
        self.assertEqual(pic.left, 0)
        self.assertAlmostEqual(pic.width / emu, pptx_export.SLIDE_W, places=2)
        self.assertAlmostEqual((pic.top + pic.height) / emu,
                               pptx_export.SLIDE_H, places=2)  # flush at the bottom
        # clear of the title and the subtitle above it
        for sh in slide.shapes:
            if sh.has_text_frame and sh.shape_id != pic.shape_id:
                self.assertLessEqual(sh.top + sh.height, pic.top + 1, sh.name)

    def test_none_and_a_missing_picture(self):
        self.assertEqual(self.pictures(self.build({"design": "none"})), [])
        notes = []
        self.assertEqual(self.pictures(self.build(
            {"design": "image", "image": "/no/such.png"}, notes)), [])
        self.assertTrue(notes)


try:
    import tkinter as tk
    import spectradeck as ee
    from readers.base import SpectrumFile, Region
    HAVE_APP = ee.HAVE_MPL and HAVE_PDF and HAVE_PPTX and HAVE_PIL
except Exception:                                    # pragma: no cover
    HAVE_APP = False


@unittest.skipUnless(HAVE_APP, "the app, matplotlib, reportlab, pptx or "
                               "Pillow missing")
class TestInTheApp(Tmp):
    @classmethod
    def setUpClass(cls):
        # a Workspace applies its theme to matplotlib's global settings;
        # give them back so other tests see the defaults
        import matplotlib
        cls._rc = matplotlib.rcParams.copy()
        try:
            cls.root = tk.Tk()
        except tk.TclError:
            raise unittest.SkipTest("no display")
        cls.root.withdraw()
        cls.tmp = tempfile.mkdtemp()
        path = os.path.join(cls.tmp, "a.vgd")
        with open(path, "wb") as fh:
            fh.write(b"stand-in")
        e = [1200 - 5.0 * i for i in range(240)]           # a survey
        r = Region(name="Survey", index=0, offset=0, energy=e,
                   counts=[10.0 + (500 if 100 < i < 110 else 0)
                           for i in range(240)],
                   decodable=True, sample="S1", photon_energy=1486.6,
                   pass_energy=160.0, dwell=0.1, step=5.0)
        f = SpectrumFile()
        f.path, f.format_name, f.regions = path, "Test", [r]
        f._finish()
        cls.doc, cls._load = f, ee.load_file
        ee.load_file = lambda p: cls.doc
        cls.ws = ee.Workspace(cls.root)
        cls.ws._add_file(path)

    @classmethod
    def tearDownClass(cls):
        ee.load_file = cls._load
        cls.root.destroy()
        import matplotlib
        matplotlib.rcParams.update(cls._rc)
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def setUp(self):
        super().setUp()
        self.ws.set_report_spec(rs.default_spec())

    def dialog(self):
        import reportgen_ui
        dlg = reportgen_ui.ReportGeneratorDialog(self.root, self.ws)
        self.addCleanup(dlg.close)
        return dlg

    def test_the_workspace_offers_a_survey_for_the_your_data_cover(self):
        energy, counts = self.ws.cover_data()
        self.assertEqual(len(energy), 240)
        self.assertEqual(len(counts), 240)

    def test_the_picker_lists_designs_with_pictures(self):
        dlg = self.dialog()
        names = [w.cget("text") for w in dlg.cover_list.winfo_children()]
        self.assertEqual(names[:6], ["No picture", "Spectrum ribbon", "Band",
                                     "Minimal", "Your data", "Peak map"])
        self.assertEqual(len(dlg._thumbs), len(names) - 1)   # none has none

    def test_choosing_a_design_a_colour_and_a_picture_changes_the_choice(self):
        dlg = self.dialog()
        dlg.set_cover(design="grid")
        self.assertEqual(rs.cover_of(self.ws.report_spec)["design"], "grid")
        dlg.set_accent("#1f7a8c")
        self.assertEqual(rs.cover_of(self.ws.report_spec)["accent"], "#1F7A8C")
        pic = self.picture("mine.png")
        dlg.choose_picture(pic)
        cover = rs.cover_of(self.ws.report_spec)
        self.assertEqual((cover["design"], cover["image"]), ("image", pic))
        names = [w.cget("text") for w in dlg.cover_list.winfo_children()]
        self.assertIn("Your picture: mine.png", names)
        self.assertEqual(dlg.design.get(), "image")

    def test_a_preset_keeps_your_cover(self):
        dlg = self.dialog()
        dlg.set_cover(design="band")
        dlg.apply_preset("Quick look")
        self.assertEqual(rs.cover_of(self.ws.report_spec)["design"], "band")
        self.assertEqual([i for i, _k in rs.active(self.ws.report_spec)],
                         ["cover", "figures"])              # the preset's sections

    def test_the_cover_reaches_the_pdf_the_deck_and_reports_a_lost_picture(self):
        ws = self.ws
        ws.set_report_spec(rs.with_cover(rs.default_spec(), design="data"))
        pdf = os.path.join(self.dir, "r.pdf")
        notes = []
        ws._build_report(pdf, notes=notes)
        with mupdf.open(pdf) as d:
            self.assertEqual(len(d[0].get_image_info()), 1)
        self.assertEqual(notes, [])
        ws.set_report_spec(rs.with_cover(rs.default_spec(), design="image",
                                         image=os.path.join(self.dir, "x.png")))
        notes = []
        ws._build_deck(os.path.join(self.dir, "d.pptx"), notes=notes)
        self.assertTrue(notes)                               # said, not silent


if __name__ == "__main__":
    unittest.main()
