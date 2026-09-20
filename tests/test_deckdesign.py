"""The redesigned PowerPoint: contents slide with real numbers, divider slides,
"n of N" footers and the cover's accent (milestone G5).

Run:  python -m unittest discover tests
"""

import io
import os
import re
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import reportspec as rs  # noqa: E402

try:
    import pptx  # noqa: F401
    import matplotlib  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


class TestOption(unittest.TestCase):
    def test_dividers_default_auto_and_bad_values_reset(self):
        self.assertEqual(rs.option(rs.default_spec(), "dividers"), "auto")
        self.assertEqual(rs.option(rs.with_option(
            rs.default_spec(), "dividers", "weird"), "dividers"), "auto")
        self.assertEqual(rs.option(rs.with_option(
            rs.default_spec(), "dividers", "none"), "dividers"), "none")

    def test_the_old_section_names_have_no_dividers(self):
        spec = rs.spec_from_sections(("title", "figures"), "deck")
        self.assertEqual(rs.option(spec, "dividers"), "none")


@unittest.skipUnless(HAVE_PPTX, "python-pptx / matplotlib not installed")
class TestDeck(unittest.TestCase):
    def setUp(self):
        from test_metasummary import Doc, region
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.docs = [Doc([region("Survey", 160, step=1.0)], "a.vgd"),
                     Doc([region("Mo 3d", 40)], "b.vgd")]
        self.rows = [{"name": "a.vgd", "format": "Thermo", "regions": 1,
                      "size": 2048, "sha256": "ab" * 32}]
        self.details = {"title": "The Study", "summary": "Everything worked.",
                        "methods": "Spectra were recorded.",
                        "calibration": "C 1s at 284.8 eV."}

    def figs(self, n):
        return [{"name": f"Fig {i}", "caption": "", "state": {}}
                for i in range(1, n + 1)]

    def png(self, number, fig):
        from matplotlib.figure import Figure
        f = Figure(figsize=(12.1, 4.95))
        f.text(0.5, 0.5, f"fig {number}")
        buf = io.BytesIO()
        f.savefig(buf, format="png", dpi=30)
        return [buf.getvalue()]

    def build(self, spec=None, figures=6):
        import pptx_export
        from pptx import Presentation
        path = os.path.join(self.dir, "d.pptx")
        self.n = pptx_export.build_deck(
            path, self.details, "", self.rows, self.docs, self.figs(figures),
            self.png, spec=spec or rs.default_spec())
        self.prs = Presentation(path)
        return list(self.prs.slides)

    @staticmethod
    def texts(slide):
        return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]

    @staticmethod
    def title(slide):
        return slide.shapes.title.text if slide.shapes.title is not None \
            else ""

    def contents_rows(self, slides):
        rows = []
        for s in slides:
            if self.title(s).startswith("Contents"):
                for sh in s.shapes:
                    if sh.has_table:
                        rows += [(r.cells[0].text.strip(), int(r.cells[1].text))
                                 for r in sh.table.rows]
        return rows

    def test_the_contents_slide_follows_the_cover_and_its_numbers_are_true(self):
        slides = self.build()
        self.assertEqual(self.title(slides[1]), "Contents")
        rows = self.contents_rows(slides)
        self.assertIn(("Data files", self.n), rows)
        for name, number in rows:
            target = slides[number - 1]
            words = self.title(target) or " ".join(self.texts(target))
            self.assertTrue(
                words.startswith(name) or name in words
                or name in self.title(target), (name, number, words))

    def test_a_divider_comes_before_a_long_section_only(self):
        slides = self.build()
        rows = dict(self.contents_rows(slides))
        fig = rows["Figures"]
        self.assertIn("Figures", " ".join(self.texts(slides[fig - 1])))
        self.assertIn("6 slides", " ".join(self.texts(slides[fig - 1])))
        self.assertTrue(self.title(slides[fig]).startswith("Figure 1"))
        # methods is one slide: no divider, the entry is the slide itself
        self.assertEqual(self.title(slides[rows["Methods"] - 1]), "Methods")

    def test_dividers_can_be_switched_off(self):
        spec = rs.with_option(rs.default_spec(), "dividers", "none")
        slides = self.build(spec)
        self.assertFalse(any("6 slides" in " ".join(self.texts(s))
                             for s in slides))
        self.assertEqual(self.title(slides[dict(
            self.contents_rows(slides))["Figures"] - 1]), "Figure 1 – Fig 1")

    def test_footers_say_n_of_n_and_the_cover_and_dividers_have_none(self):
        slides = self.build()
        self.assertNotIn(" of ", " ".join(self.texts(slides[0])))
        for k, s in enumerate(slides, 1):
            foot = [t for t in self.texts(s) if "The Study" in t
                    and " of " in t]
            divider = "slides" in " ".join(self.texts(s)) and \
                not self.title(s)
            if k == 1 or divider:
                self.assertFalse(foot, k)
            else:
                self.assertEqual(len(foot), 1, k)
                self.assertTrue(foot[0].endswith(f"{k} of {self.n}"), foot)

    def test_the_contents_can_go_last_and_the_numbers_still_hold(self):
        spec = rs.moved(rs.default_spec(), "contents", 9)
        slides = self.build(spec)
        self.assertEqual(self.title(slides[-1]), "Contents")
        rows = dict(self.contents_rows(slides))
        self.assertEqual(rows["Summary"], 2)
        self.assertEqual(self.title(slides[rows["Summary"] - 1]), "Summary")

    def test_a_long_contents_list_runs_over_and_stays_true(self):
        slides = self.build(figures=40)
        heads = [i for i, s in enumerate(slides, 1)
                 if self.title(s).startswith("Contents")]
        self.assertGreaterEqual(len(heads), 2)
        self.assertEqual(heads, list(range(heads[0], heads[0] + len(heads))))
        rows = self.contents_rows(slides)
        num = next(n for t, n in rows if t.startswith("Figure 40"))
        self.assertEqual(self.title(slides[num - 1]), "Figure 40 – Fig 40")

    def test_contents_with_nothing_to_list_makes_no_slide(self):
        spec = rs.with_all(rs.default_spec(), False)
        spec = rs.with_on(rs.with_on(spec, "cover", True), "contents", True)
        slides = self.build(spec)
        self.assertEqual(len(slides), 1)

    def test_the_accent_colours_the_bar_and_the_titles(self):
        from pptx.dml.color import RGBColor
        spec = rs.with_cover(rs.default_spec(), accent="#1F7A8C")
        slides = self.build(spec)
        summary = next(s for s in slides if self.title(s) == "Summary")
        bar = next(sh for sh in summary.shapes if sh.shape_type == 1
                   and sh.width == self.prs.slide_width)
        self.assertEqual(bar.fill.fore_color.rgb, RGBColor(0x1F, 0x7A, 0x8C))
        run = summary.shapes.title.text_frame.paragraphs[0].runs[0]
        self.assertNotEqual(run.font.color.rgb, RGBColor(0x1F, 0x7A, 0x8C))
        self.assertEqual(run.font.color.rgb, RGBColor(0x19, 0x62, 0x70))  # ink

    def test_the_old_sections_argument_is_unchanged(self):
        import pptx_export
        from pptx import Presentation
        path = os.path.join(self.dir, "o.pptx")
        n = pptx_export.build_deck(path, self.details, "", self.rows,
                                   self.docs, self.figs(6), self.png,
                                   ("title", "figures"))
        titles = [s.shapes.title.text if s.shapes.title is not None else ""
                  for s in Presentation(path).slides]
        self.assertEqual(n, 1 + 2 + 6)      # title, summary, methods, figures
        self.assertFalse(any(t.startswith("Contents") for t in titles))
        self.assertTrue(re.match(r"Figure 1", titles[3]))


if __name__ == "__main__":
    unittest.main()
