"""Quantification pages for the report and the slides (milestone G6).

Run:  python -m unittest discover tests
"""

import io
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import quant  # noqa: E402
import reportspec as rs  # noqa: E402
import resultspages as rp  # noqa: E402

try:
    import numpy  # noqa: F401
    HAVE_NP = True
except ImportError:
    HAVE_NP = False
try:
    import pymupdf as mupdf
except ImportError:
    try:
        import fitz as mupdf
    except ImportError:
        mupdf = None
try:
    import reportlab  # noqa: F401
    import matplotlib  # noqa: F401
    HAVE_PDF = mupdf is not None
except ImportError:
    HAVE_PDF = False
try:
    import pptx  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


def row(region, rsf, area, states=(), background="Shirley"):
    """A synthetic ``quant.fit_rows`` row; ``states`` is [(name, area)]."""
    comps = [{"name": n, "group": "", "index": -1, "be": 285.0, "fwhm": 1.0,
              "area": a, "shape": "GL(30)", "rsf": rsf, "gk": f"n{n}",
              "state": n} for n, a in states]
    return {"region": region, "background": background, "rsf": rsf,
            "area": area, "area_t": None, "basis": "data", "be_lo": 280.0,
            "be_hi": 290.0, "avg": 1, "rms": 0.01, "approximate": False,
            "background_known": True, "scale_known": True,
            "components": comps}


def row_src(region, rsf, area, source, **kw):
    """``row`` with a ``source`` tag ("survey" or "high-res") set."""
    r = row(region, rsf, area, **kw)
    r["source"] = source
    return r


def level(n, rows, etch=None, depth=None):
    lv = rp.Level(n, etch, depth)
    lv.entries = [{"spectrum": r["region"], "row": r} for r in rows]
    rp._settle(lv, "S", [])
    return lv


def sample(label, levels):
    return rp.Sample(f"f1/{label}", label, levels)


def three_element_level(n, etch=None, depth=None, ti=100.0):
    return level(n, [row("Ti 2p", 2.0, ti, [("Ti-C", ti * .6),
                                            ("Ti-O", ti * .4)]),
                     row("O 1s", 1.0, 100.0), row("C 1s", 0.25, 30.0)],
                 etch, depth)


class TestNumbers(unittest.TestCase):
    def test_at_percent_is_area_over_rsf_shared_out(self):
        lv = level(None, [row("A", 1.0, 100.0), row("B", 2.0, 100.0),
                          row("C", None, 50.0)])
        cells = dict((c[0], c) for _k, c in rp.composition_cells(lv))
        self.assertEqual(cells["A"][5], "66.7")
        self.assertEqual(cells["B"][5], "33.3")
        self.assertEqual(cells["C"][5], "no RSF")        # listed, with the reason

    def test_fit_rms_is_shown_as_a_percentage_and_blank_when_unknown(self):
        lv = level(None, [row("A", 1.0, 100.0),
                          dict(row("B", 2.0, 100.0), rms=None)])
        cells = dict((c[0], c) for _k, c in rp.composition_cells(lv))
        self.assertEqual(cells["A"][-1], "1.0%")            # row()'s rms=0.01
        self.assertEqual(cells["B"][-1], "")
        states = [c for k, c in rp.composition_cells(
            three_element_level(None)) if k == "state"]
        self.assertTrue(all(c[-1] == "" for c in states))

    def test_states_are_listed_under_a_region_with_more_than_one(self):
        lv = three_element_level(None)
        kinds = [k for k, _c in rp.composition_cells(lv)]
        self.assertEqual(kinds, ["region", "state", "state", "region",
                                 "region"])
        states = [c for k, c in rp.composition_cells(lv) if k == "state"]
        total = sum(float(c[5]) for c in states)
        ti = next(c for k, c in rp.composition_cells(lv) if c[0] == "Ti 2p")
        self.assertAlmostEqual(total, float(ti[5]), delta=0.11)

    def test_a_region_fitted_in_two_spectra_counts_once_the_scan_over_the_survey(
            self):
        notes = []
        lv = rp.Level(None)
        lv.entries = [{"spectrum": "Survey", "row": row("C 1s", 0.25, 30.0)},
                      {"spectrum": "C 1s", "row": row("C 1s", 0.25, 30.0)},
                      {"spectrum": "O 1s", "row": row("O 1s", 1.0, 30.0)}]
        rp._settle(lv, "S", notes)
        pct = [x["at_pct"] for x in lv.res]
        self.assertIsNone(pct[0])                     # the survey's fit
        self.assertAlmostEqual(pct[1], 100 * 120 / 150, places=6)
        self.assertEqual(lv.res[0]["why"], "counted once")
        self.assertEqual(notes, ["C 1s is fitted in more than one spectrum "
                                 "of S; only the fit in C 1s is counted."])

    def test_with_no_dedicated_scan_the_first_counts(self):
        lv = rp.Level(None)
        lv.entries = [{"spectrum": "Wide", "row": row("C 1s", 1.0, 10.0)},
                      {"spectrum": "Wide 2", "row": row("C 1s", 1.0, 30.0)}]
        rp._settle(lv, "S", [])
        self.assertEqual(lv.include, [True, False])

    def test_an_element_counted_from_two_lines_says_so(self):
        s = sample("S", [level(None, [row("Ti 2p", 5.0, 100.0),
                                      row("Ti 1s", 98.0, 900.0),
                                      row("O 1s", 3.0, 50.0),
                                      row("VB", 1.0, 5.0)])])
        rp._element_note(s)
        self.assertEqual(len(s.notes), 1)
        self.assertIn("Ti is counted from more than one line (Ti 2p, Ti 1s)",
                      s.notes[0])
        self.assertEqual(rp.element_of("Cl 2p"), "Cl")
        self.assertEqual(rp.element_of("WideScan"), "")
        self.assertEqual(rp.element_of("VB"), "")

    def test_the_same_numbers_as_quant(self):
        lv = three_element_level(None)
        direct = quant.normalise([e["row"] for e in lv.entries])
        self.assertEqual([x["at_pct"] for x in lv.res],
                         [x["at_pct"] for x in direct])

    def test_a_survey_row_is_marked_in_the_table(self):
        lv = level(None, [row_src("Cl 2p", 1.0, 10.0, "survey"),
                          row_src("C 1s", 0.25, 100.0, "high-res")])
        self.assertTrue(rp.has_survey_rows(lv))
        names = [c[0] for k, c in rp.composition_cells(lv) if k == "region"]
        self.assertEqual(names, ["Cl 2p †", "C 1s"])

    def test_no_marker_without_a_survey_row(self):
        lv = level(None, [row_src("C 1s", 0.25, 100.0, "high-res")])
        self.assertFalse(rp.has_survey_rows(lv))
        names = [c[0] for k, c in rp.composition_cells(lv) if k == "region"]
        self.assertEqual(names, ["C 1s"])

    def test_a_survey_only_element_gets_a_mixing_note(self):
        s = sample("S", [level(None, [row_src("Cl 2p", 1.0, 10.0, "survey"),
                                      row_src("C 1s", 0.25, 100.0,
                                              "high-res")])])
        rp._source_note(s)
        self.assertEqual(len(s.notes), 1)
        self.assertIn("Cl", s.notes[0])
        self.assertIn("survey scan", s.notes[0])
        self.assertIn("S", s.notes[0])

    def test_no_mixing_note_when_everything_is_one_source(self):
        s = sample("S", [level(None, [row_src("Cl 2p", 1.0, 10.0, "survey"),
                                      row_src("F 1s", 0.5, 20.0, "survey")])])
        rp._source_note(s)
        self.assertEqual(s.notes, [])
        s2 = sample("S", [level(None, [row_src("C 1s", 0.25, 100.0,
                                               "high-res"),
                                       row_src("O 1s", 1.0, 50.0,
                                               "high-res")])])
        rp._source_note(s2)
        self.assertEqual(s2.notes, [])

    def test_no_mixing_note_when_the_survey_row_is_not_counted(self):
        # the dedicated scan wins the dedup; the losing survey duplicate of
        # the same region must not itself trigger a mixing note
        lv = rp.Level(None)
        lv.entries = [{"spectrum": "Survey", "row": row_src(
                          "C 1s", 0.25, 10.0, "survey")},
                      {"spectrum": "C 1s", "row": row_src(
                          "C 1s", 0.25, 100.0, "high-res")}]
        s = sample("S", [lv])
        rp._settle(lv, "S", s.notes)
        rp._source_note(s)
        self.assertEqual([n for n in s.notes if "survey scan" in n], [])


class TestDepthProfile(unittest.TestCase):
    def profile(self, **kw):
        return sample("Film", [three_element_level(i, ti=100 - 20 * i, **kw(i))
                               for i in range(4)] if callable(kw) else [])

    def test_the_axis_is_depth_then_etch_then_level(self):
        mk = lambda **k: sample("F", [three_element_level(   # noqa: E731
            i, **{a: (v * i if v is not None else None)
                  for a, v in k.items()}) for i in range(3)])
        self.assertEqual(rp.profile_axis(mk(depth=5.0, etch=10.0))[0],
                         "Depth (nm)")
        self.assertEqual(rp.profile_axis(mk(depth=None, etch=10.0))[0],
                         "Etch time (s)")
        label, xs = rp.profile_axis(mk(depth=None, etch=None))
        self.assertEqual((label, xs), ("Level", [0, 1, 2]))

    def test_the_table_has_a_column_per_region(self):
        s = sample("F", [three_element_level(i, etch=30.0 * i,
                                             ti=100 - 20 * i)
                         for i in range(4)])
        self.assertTrue(s.is_profile)
        header, rows = rp.profile_cells(s)
        self.assertEqual(header, ["Level", "Etch time (s)", "Ti 2p", "O 1s",
                                  "C 1s"])
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[2][:2], ["2", "60"])
        # each level is normalised on its own: a row sums to 100
        self.assertAlmostEqual(sum(float(v) for v in rows[1][2:]), 100.0,
                               delta=0.2)

    def test_the_chart_is_a_png(self):
        s = sample("F", [three_element_level(i, etch=30.0 * i)
                         for i in range(4)])
        png = rp.profile_png(s, dpi=60)
        self.assertEqual(png[:8], b"\x89PNG\r\n\x1a\n")


@unittest.skipUnless(HAVE_NP, "numpy not installed")
class TestCollect(unittest.TestCase):
    def docs(self, levels=(None,)):
        from test_metasummary import Doc
        from test_quant import linear_region
        regs = []
        for lv in levels:
            r = linear_region()
            r.etch_level = lv
            r.etch_time = None if lv is None else 30.0 * lv
            regs.append(r)
        return [Doc(regs, "fits.vms")]

    def test_a_fitted_region_becomes_a_composition(self):
        res = rp.collect(self.docs())
        self.assertTrue(res)
        self.assertEqual(len(res.samples), 1)
        s = res.samples[0]
        self.assertFalse(s.is_profile)
        cells = rp.composition_cells(s.levels[0])
        self.assertEqual(cells[0][1][0], "Ti 2p")
        self.assertEqual(cells[0][1][1], "Linear")
        self.assertEqual(cells[0][1][5], "100.0")             # only region
        self.assertEqual(sum(1 for k, _c in cells if k == "state"), 2)
        self.assertEqual(res.children(), [(s.key, s.label)])

    def test_levels_make_a_depth_profile_in_order(self):
        res = rp.collect(self.docs(levels=(2, 0, 1)))
        s = res.samples[0]
        self.assertTrue(s.is_profile)
        self.assertEqual([lv.level for lv in s.levels], [0, 1, 2])
        self.assertEqual(rp.profile_axis(s)[0], "Etch time (s)")

    def test_files_without_fits_give_nothing(self):
        from test_metasummary import Doc, region
        res = rp.collect([Doc([region("Survey", 160, step=1.0)], "a.vgd")])
        self.assertFalse(res)
        self.assertEqual(res.children(), [])

    def test_the_display_copy_is_what_is_read(self):
        import copy
        seen = []

        def display(r):
            seen.append(r)
            q = copy.copy(r)
            q.sample = "Renamed"
            return q
        res = rp.collect(self.docs(), display)
        self.assertTrue(seen)
        self.assertEqual(res.samples[0].label, "Renamed")


class TestSpecAndInventory(unittest.TestCase):
    def test_the_section_sits_after_the_summary_and_is_off_for_old_names(self):
        order = rs.order(rs.default_spec())
        self.assertEqual(order[order.index("summary") + 1], "results")
        self.assertEqual(rs.LABELS["results"], "Quantification")
        self.assertNotIn("results", [i for i, _ in rs.active(
            rs.spec_from_sections(("cover", "figures"), "pdf"))])

    def test_the_customer_report_has_it_the_audit_trail_does_not(self):
        self.assertTrue(rs.is_on(rs.BUILTIN_PRESETS["Customer report"],
                                 "results"))
        self.assertFalse(rs.is_on(rs.BUILTIN_PRESETS["Audit trail"],
                                  "results"))

    def test_samples_are_children_that_can_be_switched_off(self):
        spec = rs.with_child(rs.default_spec(), "results", "f1/A", False)
        self.assertEqual(rs.skipped(spec, "results"), {"f1/A"})
        inv = rs.inventory({}, "", "", [], [], [], False,
                           results=[("f1/A", "A"), ("f1/B", "B")])
        self.assertIn("results", inv.present_ids())
        self.assertEqual(inv.summary("results"), "2 samples")
        self.assertEqual(inv.children["results"], [("f1/A", "A"),
                                                   ("f1/B", "B")])
        none = rs.inventory({}, "", "", [], [], [], False)
        self.assertNotIn("results", none.present_ids())
        self.assertIn("fits", none.summary("results"))


class Tmp(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.results = rp.Results(
            samples=[
                sample("Film A", [three_element_level(None)]),
                sample("Etched", [three_element_level(
                    i, etch=30.0 * i, ti=100 - 20 * i) for i in range(4)])])
        self.results.samples[0].notes = [
            "C 1s is fitted in more than one spectrum of Film A; "
            "the first is counted."]


@unittest.skipUnless(HAVE_PDF, "reportlab / PyMuPDF / matplotlib missing")
class TestPdf(Tmp):
    def setUp(self):
        super().setUp()
        from test_metasummary import Doc, region
        self.docs = [Doc([region("Survey", 160, step=1.0)], "a.vgd")]
        self.details = {"title": "T", "summary": "Fine.",
                        "methods": "Recorded."}

    def build(self, spec=None, results="default"):
        import report
        path = os.path.join(self.dir, "r.pdf")
        report.build_report(
            path, self.details, "", [], self.docs, [], lambda *a: 0,
            spec=spec or rs.default_spec(),
            results=self.results if results == "default" else results)
        self.doc = mupdf.open(path)
        self.addCleanup(self.doc.close)
        return [p.get_text() for p in self.doc]

    def test_the_section_has_the_numbers_the_chart_and_the_notes(self):
        text = "\n".join(self.build())
        for token in ("Quantification", "Film A", "Etched", "Ti 2p", "O 1s",
                      "Shirley", "Area / RSF", "Etch time (s)",
                      "counts/s·eV", "the first is counted",
                      "no transmission correction", "Fit RMS", "1.0%"):
            self.assertIn(token, text)
        self.assertGreaterEqual(sum(len(p.get_images())
                                    for p in self.doc), 1)   # the profile chart

    def test_it_comes_after_the_summary_and_is_in_the_contents_and_bookmarks(self):
        pages = self.build()
        toc = {t: p for _l, t, p in self.doc.get_toc()}
        self.assertIn("Quantification", toc)
        self.assertGreater(toc["Quantification"], toc["Contents"])
        self.assertIn("Film A", toc)                         # two samples: kids
        self.assertIn("Quantification", pages[toc["Quantification"] - 1])
        self.assertIn("Quantification", pages[1])            # printed contents

    def test_a_sample_can_be_left_out(self):
        spec = rs.with_child(rs.default_spec(), "results", "f1/Film A", False)
        text = "\n".join(self.build(spec))
        self.assertNotIn("Film A", text)
        self.assertIn("Etched", text)

    def test_nothing_to_quantify_leaves_the_section_out(self):
        text = "\n".join(self.build(results=None))
        self.assertNotIn("Quantification", text)

    def test_switched_off(self):
        spec = rs.with_on(rs.default_spec(), "results", False)
        self.assertNotIn("Quantification", "\n".join(self.build(spec)))


@unittest.skipUnless(HAVE_PPTX and HAVE_PDF, "python-pptx not installed")
class TestDeck(Tmp):
    def setUp(self):
        super().setUp()
        self.details = {"title": "T", "summary": "Fine."}

    def build(self, spec=None, results="default"):
        import pptx_export
        from pptx import Presentation
        path = os.path.join(self.dir, "d.pptx")
        pptx_export.build_deck(
            path, self.details, "", [], [], [], lambda n, f: [],
            spec=spec or rs.default_spec(),
            results=self.results if results == "default" else results)
        return list(Presentation(path).slides)

    @staticmethod
    def title(slide):
        return slide.shapes.title.text if slide.shapes.title is not None \
            else ""

    def test_a_table_slide_per_composition_and_chart_plus_table_per_profile(self):
        slides = self.build()
        titles = [self.title(s) for s in slides]
        self.assertIn("Quantification – Film A", titles)
        self.assertIn("Quantification – Etched: depth profile", titles)
        self.assertIn("Quantification – Etched: at % by level", titles)
        table = slides[titles.index("Quantification – Film A")]
        cells = [sh for sh in table.shapes if sh.has_table][0].table
        self.assertEqual(cells.cell(0, 5).text, "at %")
        self.assertEqual(cells.cell(0, 6).text, "Fit RMS")
        self.assertEqual(cells.cell(1, 0).text, "Ti 2p")
        self.assertEqual(cells.cell(1, 6).text, "1.0%")
        chart = slides[titles.index(
            "Quantification – Etched: depth profile")]
        self.assertTrue(any(sh.shape_type == 13 for sh in chart.shapes))
        notes = table.notes_slide.notes_text_frame.text
        self.assertIn("sensitivity factor", notes)
        self.assertIn("the first is counted", notes)

    def test_the_contents_lists_the_samples_with_true_numbers(self):
        slides = self.build()
        contents = next(s for s in slides if self.title(s) == "Contents")
        rows = [(r.cells[0].text.strip(), int(r.cells[1].text))
                for sh in contents.shapes if sh.has_table
                for r in sh.table.rows]
        got = dict(rows)
        self.assertIn("Quantification", got)
        self.assertEqual(self.title(slides[got["Film A"] - 1]),
                         "Quantification – Film A")

    def test_a_sample_can_be_left_out_or_the_section_absent(self):
        spec = rs.with_child(rs.default_spec(), "results", "f1/Film A", False)
        titles = [self.title(s) for s in self.build(spec)]
        self.assertNotIn("Quantification – Film A", titles)
        titles = [self.title(s) for s in self.build(results=None)]
        self.assertFalse(any(t.startswith("Quantification") for t in titles))


if __name__ == "__main__":
    unittest.main()
