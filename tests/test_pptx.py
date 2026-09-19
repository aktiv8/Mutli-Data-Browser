"""PowerPoint export (python-pptx). The pure helpers are always tested; the
deck tests are skipped without python-pptx.

Run:  python -m unittest discover tests
"""

import io
import math
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

import pptx_export  # noqa: E402
from test_metasummary import Doc, region  # noqa: E402


def depth(name, level, time, date, sample="D", pe=40):
    r = region(name, pe, sample=sample, date=date)
    r.etch_level, r.etch_time = level, time
    return r


class TestHelpers(unittest.TestCase):
    def test_paragraphs(self):
        self.assertEqual(pptx_export.paragraphs("a\nb\n\n\nc\n\n"),
                         ["a\nb", "c"])
        self.assertEqual(pptx_export.paragraphs(""), [])

    def test_chunk_keeps_all_text_and_respects_the_limit(self):
        paras = ["word " * 40 for _ in range(12)] + ["x" * 4000]
        slides = pptx_export.chunk_paragraphs(paras, max_lines=8,
                                              chars_per_line=90)
        joined = " ".join(" ".join(s) for s in slides)
        self.assertEqual(joined.replace(" ", "").replace("\n", ""),
                         "".join(paras).replace(" ", ""))
        for s in slides:
            est = sum(max(1, math.ceil(len(p) / 90)) + 1 for p in s)
            self.assertLessEqual(est, 8 + 1)
        self.assertGreater(len(slides), 3)

    def test_pack_items(self):
        pages = pptx_export.pack_items([(2, "a"), (2, "b"), (2, "c"),
                                        (9, "big"), (1, "d")], 5)
        self.assertEqual(pages, [["a", "b"], ["c"], ["big"], ["d"]])
        self.assertEqual(pptx_export.pack_items([], 5), [])

    def test_look_notes(self):
        text = pptx_export.look_notes({
            "view_mode": "Heatmap", "group_by": "Element name",
            "energy_scale": "Kinetic", "norm": "Max = 1",
            "colour_scale": "Viridis", "colour_reverse": True,
            "z_axis": "Auto",
            "ticked": [{"name": "C 1s"}, {"name": "C 1s"},
                       {"name": "O 1s"}]})
        for token in ("Heatmap", "kinetic energy axis", "Max = 1",
                      "Viridis (reversed)", "z axis: Auto",
                      "Spectra shown (3)", "C 1s ×2"):
            self.assertIn(token, text)
        self.assertIn("Stack", pptx_export.look_notes({}))


try:
    import pptx  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


@unittest.skipUnless(HAVE_PPTX, "python-pptx not installed")
class TestDeck(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        regs = []
        for i in range(4):
            for n, pe in (("Survey", 160), ("Mo 3d", 40), ("C 1s", 40)):
                regs.append(region(n, pe, sample=f"Sample {i}"))
        for k in range(30):
            regs.append(depth("C 1s", k, k * 15.0 if k < 20 else k * 20.0,
                              f"2020-01-01 10:{k:02d}:00"))
        self.docs = [Doc(regs, "big.avg")]
        self.rows = [{"name": "big.avg", "format": "Thermo", "regions": 42,
                      "size": 50000, "sha256": "ab" * 32}]
        self.details = {"title": "MoS2 study", "customer": "ACME",
                        "reference": "J-1", "operator": "DM", "date": "",
                        "summary": ("Long paragraph. " * 60 + "\n\n"
                                    + "Second. " * 30)}
        self.figs = [{"name": "Depth", "caption": "Caption one.\n\nTwo.",
                      "state": {"view_mode": "Heatmap"}},
                     {"name": "Stack", "caption": "",
                      "state": {"ticked": []}}]

    def png(self, number, fig):
        from matplotlib.figure import Figure
        f = Figure(figsize=pptx_export.FIGURE_SIZE)
        f.text(0.5, 0.5, f"fig {number}")
        buf = io.BytesIO()
        f.savefig(buf, format="png", dpi=60)
        return [buf.getvalue()] * (2 if number == 1 else 1)   # 2 pages

    def build(self, sections=pptx_export.SECTIONS, logo=""):
        path = os.path.join(self.tmp.name, "d.pptx")
        n = pptx_export.build_deck(path, self.details, logo, self.rows,
                                   self.docs, self.figs, self.png, sections)
        from pptx import Presentation
        return n, Presentation(path)

    @staticmethod
    def texts(slide):
        out = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    out.append(" | ".join(c.text for c in r.cells))
        return "\n".join(out)

    def test_structure_and_order(self):
        n, prs = self.build()
        slides = list(prs.slides)
        self.assertEqual(n, len(slides))
        titles = [s.shapes.title.text if s.shapes.title is not None else ""
                  for s in slides]
        self.assertEqual(titles[0], "MoS2 study")
        self.assertTrue(titles[1].startswith("Summary"))
        self.assertIn("Data files", titles)
        meta = [t for t in titles if t.startswith("Acquisition metadata")]
        self.assertGreaterEqual(len(meta), 1)
        self.assertTrue(any("per-level details" in t for t in titles))
        figs = [t for t in titles if t.startswith("Figure")]
        self.assertEqual(figs, ["Figure 1 – Depth",
                                "Figure 1 – Depth (continued)",
                                "Figure 2 – Stack"])
        order = [titles.index(t) for t in (titles[1], "Data files",
                                           meta[0], figs[0])]
        self.assertEqual(order, sorted(order))
        self.assertEqual((prs.slide_width, prs.slide_height),
                         (int(13.333 * 914400), int(7.5 * 914400)))

    def test_everything_sits_inside_the_slide(self):
        _n, prs = self.build()
        W, H = prs.slide_width, prs.slide_height
        for i, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                self.assertGreaterEqual(sh.left, 0, (i, sh.name))
                self.assertGreaterEqual(sh.top, 0, (i, sh.name))
                self.assertLessEqual(sh.left + sh.width, W + 2, (i, sh.name))
                self.assertLessEqual(sh.top + sh.height, H + 2, (i, sh.name))

    def test_text_fits_its_box(self):
        """A rough guard: estimated wrapped height <= the box height."""
        _n, prs = self.build()
        for i, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                if not sh.has_text_frame or sh == slide.shapes.title:
                    continue
                width_in = sh.width / 914400
                need = 0.0
                for p in sh.text_frame.paragraphs:
                    size = next((r.font.size.pt for r in p.runs
                                 if r.font.size), 14)
                    per_line = max(10, width_in * 72 / (size * 0.5))
                    lines = max(1, math.ceil(len(p.text) / per_line))
                    after = p.space_after.pt if p.space_after else 0
                    need += lines * size * 1.2 + after
                self.assertLessEqual(need / 72, sh.height / 914400 + 0.35,
                                     (i, sh.text_frame.text[:40]))

    def test_tables_and_lossless_metadata_on_slides(self):
        _n, prs = self.build(("metadata",))
        text = "\n".join(self.texts(s) for s in prs.slides)
        for token in ("Common to every region", "Sample 0", "Sample 3",
                      "Survey", "Mo 3d", "160", "0–29 (30)",
                      "Level", "Etch (s)"):
            self.assertIn(token, text)
        # the irregular etch times (k*20 for k >= 20) are all spelled out
        for k in range(30):
            self.assertIn(f"{k * 15.0 if k < 20 else k * 20.0:g}", text)

    def test_figure_slides_have_picture_caption_and_notes(self):
        _n, prs = self.build(("figures",))
        first = list(prs.slides)[0]
        kinds = [sh.shape_type for sh in first.shapes]
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        self.assertIn(MSO_SHAPE_TYPE.PICTURE, kinds)
        self.assertIn("Caption one.", self.texts(first))
        notes = first.notes_slide.notes_text_frame.text
        self.assertIn("Caption one.", notes)
        self.assertIn("Heatmap", notes)

    def test_sections_are_optional_and_empty_is_an_error(self):
        n, prs = self.build(("title",))
        self.assertGreaterEqual(n, 2)
        with self.assertRaises(pptx_export.PptxError):
            self.figs = []
            self.build(("figures",))

    def test_logo_is_placed_and_a_bad_one_ignored(self):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("Pillow not installed")
        logo = os.path.join(self.tmp.name, "logo.png")
        Image.new("RGB", (600, 300), (20, 90, 160)).save(logo)
        _n, prs = self.build(("title",), logo)
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        pics = [sh for sh in list(prs.slides)[0].shapes
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pics), 1)
        self.assertLessEqual(pics[0].height, int(1.3 * 914400) + 1)
        bad = os.path.join(self.tmp.name, "bad.png")
        with open(bad, "wb") as fh:
            fh.write(b"nope")
        self.build(("title",), bad)             # must not raise


if __name__ == "__main__":
    unittest.main()
