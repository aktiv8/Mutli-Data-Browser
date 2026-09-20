"""Camera pictures and SnapMaps as report pages and slides (imagepages.py) and
their place in the PDF report and the deck. Uses the synthetic maps and
pictures of test_htmlbrowser; needs numpy, Pillow and matplotlib."""

import copy
import io
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

try:
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    import numpy as np
    import PIL  # noqa: F401
    HAVE = True
except Exception:                                    # pragma: no cover
    HAVE = False

if HAVE:
    import imagepages as ip
    import pptx_export
    import report
    from readers import ImageBlob
    from readers.imaging import encode_png
    from test_htmlbrowser import (CAMERA_CAL, camera_blob, cube_region, doc,
                                  map_cube, region)


def two_pictures(n=2):
    """``n`` camera pictures at the map's position (times differ)."""
    return [camera_blob(f"S1 Pt #001a  1{i}:00") for i in range(n)]


def site_doc(n_pictures=1, with_map=True):
    regs = [region("C 1s", "S2")]
    if with_map:
        regs.append(cube_region(map_cube()))
    return doc("a.vgd", regs, positions={"S1": (10.0, 20.0), "S2": (10.1, 20.1)},
               images=two_pictures(n_pictures) if n_pictures else [])


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestPlan(unittest.TestCase):
    def test_sheets_come_before_map_pages(self):
        pages = ip.plan([site_doc(2)])
        self.assertEqual([(p.kind, p.title, p.n_items) for p in pages],
                         [("camera", "Camera pictures", 2),
                          ("maps", "SnapMap – S1", 1)])
        self.assertTrue(ip.available([site_doc(2)]))

    def test_long_runs_are_split_into_numbered_sheets(self):
        pages = ip.plan([site_doc(7, with_map=False)], per_sheet=3)
        self.assertEqual([(p.title, p.n_items) for p in pages],
                         [("Camera pictures (1 of 3)", 3),
                          ("Camera pictures (2 of 3)", 3),
                          ("Camera pictures (3 of 3)", 1)])

    def test_only_calibrated_pictures_and_real_maps_count(self):
        plain = ImageBlob("holder", 0, b"x", True)
        d = doc("a.vms", [region("C 1s", "S1")], images=[plain])
        self.assertEqual(ip.plan([d]), [])
        self.assertFalse(ip.available([d]))
        self.assertTrue(ip.available([site_doc(0)]))          # a map alone
        self.assertEqual([p.kind for p in ip.plan([site_doc(0)])], ["maps"])

    def test_a_map_finds_the_picture_taken_at_its_position(self):
        pages = ip.plan([site_doc(1)])
        # planned internally: the map page's notes name the camera picture
        self.assertIn("Camera picture: S1 Pt #001a", pages[1].notes())
        self.assertNotIn("Camera picture", ip.plan([site_doc(0)])[0].notes())

    def test_names_and_energies_are_shown_as_the_app_shows_them(self):
        def display(r):
            q = copy.copy(r)
            q.name = "Renamed " + r.name
            q.energy = [e + 1.5 for e in r.energy]
            return q
        pages = ip.plan([site_doc(1)], label_of=lambda p, s: "Lab " + s,
                        display=display)
        self.assertEqual(pages[1].title, "SnapMap – Lab S1")
        note = pages[1].notes()
        self.assertIn("Renamed C 1s: counts summed over", note)
        cube = map_cube()
        lo, hi = ip.window_of(display(cube_region(cube)), cube)
        import snapmap
        w = snapmap.default_window(cube)
        self.assertAlmostEqual(lo, w[0] + 1.5)
        self.assertAlmostEqual(hi, w[1] + 1.5)
        # and the picture is labelled with the shown names too
        fig = Figure(figsize=(11.7, 8.3))
        pages[0].draw(fig)
        texts = [t.get_text() for t in fig.axes[0].texts]
        self.assertIn("Lab S1", texts)                 # the map's outline
        self.assertIn("Lab S2", texts)                 # another analysis point

    def test_notes_describe_the_pictures(self):
        note = ip.plan([site_doc(2)])[0].notes()
        self.assertEqual(len(note.splitlines()), 2)
        self.assertIn("stage 10.000, 20.000 mm", note)
        self.assertIn("field of view 3.2 x 2.4 mm", note)


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestDrawing(unittest.TestCase):
    def draw(self, page, size=(11.7, 8.3), rect=(0.0, 0.05, 1.0, 0.93)):
        fig = Figure(figsize=size, dpi=50)
        page.draw(fig, rect)
        return fig

    def test_a_sheet_has_one_axes_per_picture_of_a_fixed_size(self):
        full = self.draw(ip.plan([site_doc(3, False)], per_sheet=3)[0])
        part = self.draw(ip.plan([site_doc(1, False)], per_sheet=3)[0])
        self.assertEqual(len(full.axes), 3)
        self.assertEqual(len(part.axes), 1)
        a, b = full.axes[0].get_position(), part.axes[0].get_position()
        self.assertAlmostEqual(a.width, b.width)       # same cell size

    def test_a_map_page_has_picture_map_and_colour_bar(self):
        fig = self.draw(ip.plan([site_doc(1)])[1])
        self.assertEqual(len(fig.axes), 3)             # picture, map, colour bar
        fig = self.draw(ip.plan([site_doc(0)])[0])
        self.assertEqual(len(fig.axes), 2)             # map, colour bar
        titles = [a.get_title() for a in fig.axes]
        self.assertTrue(any(t.startswith("C 1s") for t in titles))

    def test_drawing_stays_inside_the_rectangle_it_was_given(self):
        rect = (0.0, 0.05, 1.0, 0.93)
        for page in ip.plan([site_doc(2)]):
            fig = self.draw(page, rect=rect)
            for ax in fig.axes:
                pos = ax.get_position()
                self.assertGreaterEqual(pos.x0, rect[0] - 1e-9)
                self.assertLessEqual(pos.x1, rect[2] + 1e-9)
                self.assertGreaterEqual(pos.y0, rect[1] - 1e-9)
                self.assertLessEqual(pos.y1, rect[3] + 1e-9)

    def test_slide_sized_pages_draw_too(self):
        for page in ip.plan([site_doc(3)], per_sheet=3):
            fig = self.draw(page, size=pptx_export.FIGURE_SIZE,
                            rect=(0.0, 0.0, 1.0, 1.0))
            buf = io.BytesIO()
            fig.savefig(buf, format="png")
            self.assertTrue(buf.getvalue().startswith(b"\x89PNG"))

    def test_pictures_are_shrunk_for_the_page(self):
        w, h = 1600, 1200
        blob = ImageBlob("big", 0, encode_png(w, h, bytes(w * h * 3)), False,
                         fmt="png", sample="S1",
                         calib=dict(CAMERA_CAL, width=w, height=h,
                                    um_per_px_x=5.0, um_per_px_y=5.0))
        d = doc("a.vgd", [region("C 1s", "S1")], positions={"S1": (10.0, 20.0)},
                images=[blob])
        page = ip.plan([d])[0]
        fig = self.draw(page)
        img = fig.axes[0].images[0].get_array()
        self.assertEqual(img.shape[1], ip.PICTURE_PX)
        # ... but the axes still speak in the full picture's pixels
        self.assertEqual(fig.axes[0].get_xlim(), (0, w))

    def test_a_map_outline_replaces_the_marker_of_a_map_site(self):
        fig = self.draw(ip.plan([site_doc(1, True)], per_sheet=1)[0])
        ax = fig.axes[0]
        texts = [t.get_text() for t in ax.texts]
        self.assertIn("S1", texts)                     # the outline's label
        self.assertEqual(len(ax.patches) >= 1, True)
        # S1 is a map site: no second, round marker for it
        self.assertEqual(sum(1 for t in texts if t == "S1"), 1)


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestInReportAndDeck(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.details = {"title": "Study"}

    def render_images(self, pdf):
        n = 0
        for pg in ip.plan([site_doc(2)]):
            fig = Figure(figsize=(11.7, 8.3))
            pg.draw(fig, (0.0, 0.03, 1.0, 0.93))
            fig.text(0.03, 0.975, pg.title)
            pdf.savefig(fig)
            n += 1
        return n

    def render_figure(self, pdf, number, figure):
        fig = Figure(figsize=(11.7, 8.3))
        fig.text(0.5, 0.5, f"FIGURE {number}")
        pdf.savefig(fig)
        return 1

    def pages(self, path):
        try:
            import pymupdf as mu
        except ImportError:
            import fitz as mu
        with mu.open(path) as d:
            return [p.get_text() for p in d]

    def build(self, sections, render_images="default"):
        path = os.path.join(self.tmp, "r.pdf")
        ri = self.render_images if render_images == "default" else render_images
        n = report.build_report(path, self.details, "", [], [],
                                [{"name": "F", "caption": "", "state": {}}],
                                self.render_figure, sections, ri)
        return n, self.pages(path)

    def test_images_sit_between_metadata_and_figures(self):
        self.assertEqual(report.SECTIONS,
                         ("cover", "metadata", "images", "figures"))
        n, text = self.build(("cover", "images", "figures"))
        self.assertEqual(n, len(text))
        self.assertEqual(n, 4)                         # cover, sheet, map, figure
        self.assertIn("Camera pictures", text[1])
        self.assertIn("SnapMap", text[2])
        self.assertIn("FIGURE 1", text[3])
        self.assertIn("report page 2 of 4", text[1])   # numbered with the rest

    def test_images_are_left_out_when_not_chosen_or_not_offered(self):
        n, text = self.build(("cover", "figures"))
        self.assertEqual(n, 2)
        n, text = self.build(("cover", "images", "figures"), render_images=None)
        self.assertEqual(n, 2)
        with self.assertRaises(report.ReportError):
            self.build(("images",), render_images=None)

    def test_deck_has_a_slide_per_page_before_the_figures(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        def pages():
            out = []
            for pg in ip.plan([site_doc(2)], per_sheet=3):
                fig = Figure(figsize=pptx_export.FIGURE_SIZE, dpi=60)
                pg.draw(fig, (0, 0, 1, 1))
                buf = io.BytesIO()
                fig.savefig(buf, format="png")
                out.append({"title": pg.title, "png": buf.getvalue(),
                            "notes": pg.notes()})
            return out

        def fig_png(number, fig):
            f = Figure(figsize=pptx_export.FIGURE_SIZE, dpi=40)
            buf = io.BytesIO()
            f.savefig(buf, format="png")
            return [buf.getvalue()]

        path = os.path.join(self.tmp, "d.pptx")
        figs = [{"name": "F", "caption": "", "state": {}}]
        n = pptx_export.build_deck(path, self.details, "", [], [], figs, fig_png,
                                   ("title", "images", "figures"), pages)
        prs = Presentation(path)
        slides = list(prs.slides)
        self.assertEqual(n, len(slides))
        titles = [s.shapes.title.text for s in slides]
        self.assertEqual(titles[-3:], ["Camera pictures", "SnapMap – S1",
                                       "Figure 1 – F"])
        cam = slides[-3]
        self.assertIn(MSO_SHAPE_TYPE.PICTURE, [sh.shape_type for sh in cam.shapes])
        self.assertIn("stage 10.000, 20.000 mm",
                      cam.notes_slide.notes_text_frame.text)
        # nothing to add: no callable, no slides
        n2 = pptx_export.build_deck(path, self.details, "", [], [], figs, fig_png,
                                    ("title", "images", "figures"), None)
        self.assertEqual(n2, n - 2)

    def test_the_deck_default_sections_include_images(self):
        self.assertEqual(pptx_export.SECTIONS,
                         ("title", "files", "metadata", "images", "figures"))


if __name__ == "__main__":
    unittest.main()
